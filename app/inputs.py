"""Reads the two input files straight into the solver's input (core.model.ScheduleInput).

COMBO APPROVALS (Combo Approvals.xlsx, table Approvals, filled by the approval flow; the director decides Status).
Columns found by words in the header, so exact wording and order don't matter:
    Response Id | Submitted | Semester | Liaison | Members | Supervisor | First year | Status | Decided by | ... | Notes
      - Members: every member's email, any separator and any text around them ("piano: ana.ruiz@..." works);
        the liaison is added if missing; domain slips are corrected (settings email_domain_fixes, e.g. a bare
        @mcgill.ca becomes @mail.mcgill.ca)
      - Status: Pending, Accepted, Rejected or Withdrawn (any capitalisation). Only Accepted rows are scheduled;
        Pending ones are reported (level "pending"), anything else gets a warning
      - optional: a header with "combo" and "name" -> that combo's name instead of Combo 01, 02, ...
    Corrected emails (the app's Combos tab, kept in scheduler_data.json) replace the typed ones here and in the
    conflicts, so a typo in either file can be fixed without touching it.
    Only rows whose Semester matches the settings' semester_name count. The same members accepted twice: the later
    row wins. Combo numbers follow Response Id order and are kept in scheduler_data.json (store.py), so a combo never
    changes number; a combo that stops being accepted leaves a gap.
CONFLICTS (Conflicts.xlsx, table Conflicts, filled by the conflict flow; the raw Forms export also works):
    Response ID | Submitted | Semester | Email | Date 1 | ... | Date 4 | Reason | Additional Info | Status | Notes
      - Reason, Additional Info and Notes are for people; only the columns below are read
      - every header containing "date" is a date column; with two email columns (the Forms export), the typed one
        is the student and the login Email is the fallback (unless it's "anonymous")
      - Status: Active (the default: it counts) or Overruled (the director decided it doesn't). Blank, OK, Accepted
        and Counted also count; Rejected and Ignored don't. Anything else counts, with a warning.
Rules: only rows whose Semester matches Settings semester_name count. If someone submits twice, the latest row wins.
Students submit the conflict form only if they have conflicts: no submission = free every night.
Dates are month-first (11/3/2026 = Nov 3), real dates and Excel date numbers also work.

This is the only file that knows what the inputs look like. If data collection changes, replace it with something
that returns the same ScheduleInput.
"""
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from core.model import Combo, ScheduleInput
from store import Store
from util import APPROVALS_FILE, CONFLICTS_FILE, blank, to_date  # noqa: F401  (others import the names from here)

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+'\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
EXAMPLE_NAMES = {"first.last"}                      # placeholder addresses from the form text (first.last@<domain>)

# Used only if the conflicts export has no header row (the first cell is a number instead of a title).
DEFAULT_CONFLICT_HEADERS = ["Id", "Start time", "Completion time", "Email", "Name", "Semester", "Your McGill Email",
                            "Conflicting Date", "Conflicting Date1", "Conflicting Date2", "Conflicting Date3",
                            "Additional information about your scheduling conflicts (optional)"]


class InputError(Exception):
    pass


class EmailRules:
    """The settings' email rules. norm(): lowercase and correct a domain slip (email_domain_fixes, e.g. mcgill.ca ->
    mail.mcgill.ca); any other address (e.g. Gmail for a member from elsewhere) is kept as typed. is_student(): on
    student_email_domain (always True when that's blank)."""

    def __init__(self, domain="", fixes=None):
        self.domain, self.fixes = domain, fixes or {}

    @classmethod
    def from_settings(cls, settings):
        return cls(settings.student_email_domain, settings.email_domain_fixes)

    def norm(self, e):
        e = str(e or "").strip().lower()
        user, at, dom = e.rpartition("@")
        return f"{user}@{self.fixes[dom]}" if at and dom in self.fixes else e

    def is_student(self, e):
        return not self.domain or e.endswith("@" + self.domain)

    def is_example(self, e):
        return e.rpartition("@")[0] in EXAMPLE_NAMES


def same(a, b):
    return str(a or "").strip().casefold() == str(b or "").strip().casefold()


def first_email(rules, *cells):
    for c in cells:
        m = EMAIL_RE.search(str(c or ""))
        if m:
            return rules.norm(m.group(0))
    return ""


def read_export(path, sheet, defaults, what, notes, needs=()):
    """Returns (headers, rows) with rows = [(excel_row, order_key, values)]. order_key = Id / Response Id (or row
    number). Without a sheet name, uses the first sheet whose header row contains all the words in `needs`."""
    try:
        wb = load_workbook(path, data_only=True)
    except FileNotFoundError:
        raise InputError(f"Can't find {path}.")
    except PermissionError:
        raise InputError(f"Can't open {path}. Close it in Excel and run again.")
    except (zipfile.BadZipFile, KeyError, InvalidFileException, OSError):
        # e.g. caught half-written while OneDrive syncs a new version (the flow just added a row)
        raise InputError(f"Couldn't read {Path(path).name}: it may be in the middle of syncing, or it isn't a real "
                         "Excel file. Wait a moment and try again (Reload).")
    if sheet and sheet not in wb.sheetnames:
        raise InputError(f"{what}: {path} has no sheet named '{sheet}'. Sheets: {', '.join(wb.sheetnames)}")
    ws = wb[sheet] if sheet else next(
        (w for w in wb.worksheets
         if all(any(k in str(c or "").lower() for c in next(w.iter_rows(max_row=1, values_only=True), ())) for k in needs)),
        wb.worksheets[0])
    raw = [(i, r) for i, r in enumerate(ws.iter_rows(values_only=True), start=1) if any(not blank(c) for c in r)]
    if not raw:
        raise InputError(f"{what}: {path} has no data.")
    first = raw[0][1][0]
    if defaults and (isinstance(first, (int, float)) or str(first).strip().isdigit()):
        notes.append(("info", f"{what}: no header row found; assuming the standard column order."))
        width = max(len(r) for _, r in raw)
        headers = (defaults + [""] * width)[:max(width, len(defaults))]
        data = raw
    else:
        headers = [str(h).strip() if h is not None else "" for h in raw[0][1]]
        data = raw[1:]
    id_i = next((i for i, h in enumerate(headers) if h.lower() in ("id", "response id")), 0)
    out = []
    for xl, r in data:
        try:
            key = float(r[id_i])
        except (TypeError, ValueError, IndexError):
            key = float(xl)
        out.append((xl, key, r))
    out.sort(key=lambda t: t[1])
    return headers, out


def col(headers, *keys, exact=False):
    for i, h in enumerate(headers):
        hl = h.lower()
        if (exact and hl in keys) or (not exact and any(k in hl for k in keys)):
            return i
    return None


def cell(row, i):
    return row[i] if i is not None and i < len(row) else None


def semester_filter(rows, sem_i, semester, what, notes):
    keep, other, blank_sem = [], 0, 0
    for xl, key, r in rows:
        v = cell(r, sem_i)
        if same(v, semester):
            keep.append((xl, key, r))
        elif blank(v):
            blank_sem += 1
        else:
            other += 1
    if other:
        notes.append(("info", f"{what}: skipped {other} row(s) from other semesters."))
    if blank_sem:
        notes.append(("info", f"{what}: skipped {blank_sem} row(s) with no semester (test submissions?)."))
    return keep


COUNTS = ("", "active", "ok", "okay", "accepted", "counted", "counts")
OVERRULED = ("overruled", "rejected", "ignored", "not counted")


def parse_conflicts(path, sheet, semester, notes, fix=lambda e: e, rules=EmailRules()):
    what = "Conflicts"
    headers, rows = read_export(path, sheet, DEFAULT_CONFLICT_HEADERS, what, notes)
    sem_i = col(headers, "semester")
    login_i = col(headers, "email", exact=True)
    typed_i = next((i for i, h in enumerate(headers) if "email" in h.lower() and i != login_i), None)
    date_is = [i for i, h in enumerate(headers) if "date" in h.lower()]
    status_i = col(headers, "status")
    if sem_i is None or not date_is or (login_i is None and typed_i is None):
        raise InputError(f"Conflicts export: couldn't find the Semester, email and date columns.\n  Headers seen: {headers}")
    rows = semester_filter(rows, sem_i, semester, what, notes)

    best, submissions = {}, 0
    for xl, key, r in rows:
        login = cell(r, login_i)
        email = fix(first_email(rules, cell(r, typed_i), "" if same(login, "anonymous") else login))
        if not email:
            notes.append(("warn", f"Conflicts row {xl}: no email; skipped."))
            continue
        submissions += 1
        best[email] = (xl, r)                  # rows are sorted by Id, so the latest replaces earlier ones
    if submissions > len(best):
        notes.append(("info", f"Conflicts: {submissions - len(best)} repeat submission(s); each person's latest was used."))

    blocked, overruled = {}, []
    for email, (xl, r) in sorted(best.items()):
        status = str(cell(r, status_i) or "").strip().lower()
        if status in OVERRULED:
            overruled.append(email)
            continue
        if status not in COUNTS:
            notes.append(("warn", f"Conflicts row {xl} ({email}): Status '{cell(r, status_i)}' isn't Active or "
                                  "Overruled; counted as Active."))
        ds = set()
        for i in date_is:
            try:
                d = to_date(cell(r, i))
            except ValueError as e:
                notes.append(("warn", f"Conflicts row {xl} ({email}): {e}; ignored."))
                continue
            if d:
                ds.add(d)
        blocked[email] = ds
    if overruled:
        notes.append(("info", f"Conflicts: {len(overruled)} submission(s) overruled by the director, not counted: "
                              f"{', '.join(overruled)}."))
    empty = sorted(e for e, ds in blocked.items() if not ds)
    if empty:
        notes.append(("info", f"Conflicts: {len(empty)} submission(s) listed no dates (e.g. {', '.join(empty[:3])}); "
                              "nothing to block for them."))
    blocked = {e: ds for e, ds in blocked.items() if ds}
    notes.append(("info", f"Conflicts: {len(blocked)} students have conflicts. Everyone else is treated as free."))
    return blocked


ACCEPTED, REJECTED, WITHDRAWN = ("accepted",), ("rejected",), ("withdrawn",)


def parse_approvals(path, sheet, semester, notes, store, rules=EmailRules(), use_first_year=True):
    """The approvals table -> accepted Combos, numbered with (and adding to) store's numbers for this semester."""
    what = "Approvals"
    headers, rows = read_export(path, sheet, None, what, notes, needs=("status", "members"))
    sem_i, members_i, status_i = col(headers, "semester"), col(headers, "member"), col(headers, "status")
    if sem_i is None or members_i is None or status_i is None:
        raise InputError(f"{what}: couldn't find the Semester, Members and Status columns in {Path(path).name}.\n"
                         f"  Headers seen: {headers}")
    liaison_i = col(headers, "liaison")
    prof_i = col(headers, "supervisor", "professor")
    fy_i = col(headers, "first year", "first-year")
    name_i = next((i for i, h in enumerate(headers) if "combo" in h.lower() and "name" in h.lower()), None)
    id_i = next((i for i, h in enumerate(headers) if h.lower() in ("id", "response id")), None)
    if id_i is None:
        notes.append(("warn", f"{what}: no 'Response Id' column; combos are numbered by row order instead."))
    seen_sems = sorted({str(cell(r, sem_i)).strip() for _, _, r in rows if not blank(cell(r, sem_i))})
    rows = semester_filter(rows, sem_i, semester, what, notes)
    if not rows:
        notes.append(("warn", f"{what}: no submissions for {semester}. Semesters in the table: "
                              f"{', '.join(seen_sems) or 'none'}. Is semester_name in the settings right?"))

    found, pending, rejected, withdrawn, seen_refs, outside, edited = [], [], [], [], set(), set(), {}
    for xl, key, r in rows:
        status = str(cell(r, status_i) or "").strip().lower()
        ref = ref_of(cell(r, id_i)) if id_i is not None and not blank(cell(r, id_i)) else f"row{xl}"
        seen_refs.add(ref)
        liaison = store.fix_email(first_email(rules, cell(r, liaison_i)))
        who = f"response {ref} ({liaison or 'no liaison'})"
        if status in REJECTED:
            rejected.append(who)
            continue
        if status in WITHDRAWN:
            withdrawn.append(who)
            continue
        if status in ("", "pending"):
            pending.append(who)
            continue
        if status not in ACCEPTED:
            notes.append(("warn", f"{what} {who}: Status '{cell(r, status_i)}' isn't Pending, Accepted, Rejected or "
                                  "Withdrawn; left out."))
            continue
        members, seen = [], set()
        for m in EMAIL_RE.findall(" ".join(str(cell(r, i) or "") for i in (liaison_i, members_i))):
            e = store.fix_email(rules.norm(m))
            if rules.is_example(e):
                notes.append(("warn", f"{what} {who}: example address '{e}' found; removed."))
            elif e not in seen:
                seen.add(e)
                members.append(e)
                if not rules.is_student(e):
                    outside.add(e)
        # members added or removed in the app (Combos tab) count as if the approvals said so
        changes, had_members = store.member_changes(semester, ref), bool(members)
        for e in changes["remove"]:
            if e in seen:
                members.remove(e)
                seen.discard(e)
            else:
                notes.append(("info", f"{what} {who}: {e} was removed in the app but isn't in the approvals any "
                                      "more; nothing to do."))
        for e in changes["add"]:
            if e in seen:
                notes.append(("info", f"{what} {who}: {e} was added in the app and is now in the approvals too."))
            else:
                members.append(e)
                seen.add(e)
                if not rules.is_student(e):
                    outside.add(e)
        if changes["add"] or changes["remove"] or changes["liaison"]:
            edited[ref] = changes
        if not members and not had_members:
            notes.append(("warn", f"{what} {who}: accepted but has no valid member emails; skipped."))
            continue
        if changes["liaison"] in seen:                  # chosen in the app
            liaison = changes["liaison"]
        elif liaison not in seen:                       # e.g. the liaison was removed: the first member stands in
            liaison = members[0] if members else ""
        prof_m = EMAIL_RE.search(str(cell(r, prof_i) or "")) if prof_i is not None else None
        prof = store.fix_email(prof_m.group(0).lower()) if prof_m else ""
        if prof_i is not None and not prof:
            notes.append(("warn", f"{what} {who}: no supervisor email."))
        elif rules.domain and rules.is_student(prof):
            notes.append(("warn", f"{what} {who}: supervisor '{prof}' is a student address. A mistake?"))
        found.append(dict(xl=xl, key=key, ref=ref, members=frozenset(members), liaison=liaison, prof=prof,
                          name=str(cell(r, name_i)).strip() if name_i is not None and not blank(cell(r, name_i)) else "",
                          fy=use_first_year and fy_i is not None
                          and str(cell(r, fy_i) or "").strip().lower() in ("yes", "y", "true", "1")))

    if pending:
        notes.append(("pending", f"{len(pending)} combo(s) are still Pending (no decision yet), so they are NOT "
                                 f"scheduled: {', '.join(pending)}."))
    outside &= {e for f in found for e in f["members"]}
    if outside:
        notes.append(("info", f"{what}: {len(outside)} member(s) with an email outside {rules.domain}: "
                              f"{', '.join(sorted(outside))}. "
                              "That's fine; their conflicts count if they submit the conflict form with this same "
                              "address."))
    if rejected:
        notes.append(("info", f"{what}: {len(rejected)} rejected combo(s) left out."))
    if withdrawn:
        notes.append(("info", f"{what}: {len(withdrawn)} withdrawn combo(s) left out: {', '.join(withdrawn)}."))

    latest = {}                                    # the same combo accepted twice (identical members): latest wins
    for f in found:
        if f["members"] in latest:
            notes.append(("warn", f"{what}: responses {latest[f['members']]['ref']} and {f['ref']} are both accepted "
                                  "with identical members; using the later one. Set the other to Rejected."))
        latest[f["members"]] = f
    kept = sorted(latest.values(), key=lambda f: f["key"])

    # numbers: the ones given out before stay; new combos continue after the highest (also past withdrawn ones)
    numbers = store.numbers(semester)
    nxt = max(numbers.values(), default=0) + 1
    combos = []
    for f in kept:
        if f["ref"] not in numbers:
            store.set_number(semester, f["ref"], nxt)
            nxt += 1
    digits = max(2, len(str(max(numbers.values(), default=1))))
    for f in kept:
        name = f["name"] or f"Combo {numbers[f['ref']]:0{digits}d}"
        combos.append(Combo(id=name, name=name, members=f["members"], first_year=f["fy"], liaison=f["liaison"],
                            professor=f["prof"], ref=f["ref"]))
    kept_refs = {f["ref"] for f in kept}
    gone = sorted(f"Combo {n:0{digits}d}" for ref, n in numbers.items() if ref in seen_refs and ref not in kept_refs)
    if gone:
        notes.append(("warn", f"No longer accepted: {', '.join(gone)}. If the schedule is already out, give their "
                              "sets away or set them to OPEN (Swaps or Schedule tab); check with step 3."))
    for c in combos:
        if c.ref in edited:
            ch = edited[c.ref]
            what_changed = ([f"added {e}" for e in ch["add"]] + [f"removed {e}" for e in ch["remove"]]
                            + ([f"liaison {ch['liaison']}"] if ch["liaison"] else []))
            notes.append(("info", f"{c.name}: changed in the app (not in the approvals): {', '.join(what_changed)}."))
    names = [c.name for c in combos]
    for n in sorted({n for n in names if names.count(n) > 1}):
        notes.append(("warn", f"Two combos are both named '{n}'. Rename one; the schedule can't tell them apart."))
    for i, a in enumerate(combos):               # near-duplicates: the same combo resubmitted, both accepted?
        for b in combos[i + 1:]:
            shared = len(a.members & b.members)
            if shared >= 3 and shared >= 0.6 * min(len(a.members), len(b.members)):
                notes.append(("warn", f"{a.name} and {b.name} share {shared} members: the same combo accepted twice? "
                                      "If so, set the older one to Rejected."))
    notes.append(("info", f"Combos: {len(combos)} accepted for {semester}."))
    return combos


def name_from_email(email):
    """eloise.ferreira2@mail.mcgill.ca -> 'Eloise Ferreira'; jean-luc.o'brien -> "Jean-Luc O'Brien"."""
    local = re.sub(r"\d+$", "", str(email).split("@")[0])
    words = [w for w in re.split(r"[._]+", local) if w]

    def cap(w):
        return "-".join("'".join(p[:1].upper() + p[1:] for p in h.split("'")) for h in w.split("-"))
    return " ".join(cap(w) for w in words) or str(email)


def ref_of(v):
    """Response Ids may come back as 14, 14.0 or '14'."""
    s = str(v).strip()
    return s[:-2] if s.endswith(".0") else s


def name_of_fn(store):
    """email -> the name to show: the corrected one (Combos tab) or one guessed from the email."""
    return lambda e: store.names.get(e) or name_from_email(e)


def load_input(folder, settings, approvals=APPROVALS_FILE, conflicts=CONFLICTS_FILE, approvals_sheet=None,
               conflicts_sheet=None, store=None):
    """The data folder's two input files -> ScheduleInput (combos, conflicts, every note as (level, text)).
    New combo numbers are saved to scheduler_data.json right away, so they never change. Raises InputError."""
    folder = Path(folder)
    try:
        store = store or Store(folder)
    except ValueError as e:
        raise InputError(str(e))
    notes = []
    rules = EmailRules.from_settings(settings)
    blocked = parse_conflicts(folder / conflicts, conflicts_sheet, settings.semester_name, notes, store.fix_email, rules)
    combos = parse_approvals(folder / approvals, approvals_sheet, settings.semester_name, notes, store, rules,
                             settings.use_first_year)
    if store.changed:
        store.save()
    return ScheduleInput(combos=combos, blocked=blocked, notes=notes)
