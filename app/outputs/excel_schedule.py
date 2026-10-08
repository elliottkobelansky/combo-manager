"""Schedule.xlsx: an export of the schedule (schedule_file.py) to read, print, sort or share. Rebuilt after every
change and never read back, so editing it changes nothing.

    By night    one table per show night: set, time, combo, each member (liaison marked) with their instrument
                (in the Combos tab's order), and the coach
    All sets    one row per set (date, venue, set, times, combo, FB = feedback night): for sorting and filtering
    Feedback nights  the feedback nights and who plays them (when every combo has one)
    Changes     every change since the schedule was made (swaps, give-aways, text in sets, withdrawn combos), newest
                first: when, on which computer, what, and each set's before and after
    Report      what the solver said when the schedule was made

At the end, the readers for the old, hand-editable Schedule.xlsx (before 2026-10-06), used once to turn one into
schedule.json (schedule_file.convert_old).
"""
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side

from core.model import WEEKDAYS, make_label
from schedule_file import ScheduleFileError
from shared_folder import save_workbook
from util import by_instrument

FONT = "Arial"
HEAD_FILL, NIGHT_FILL = PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="B4C6E7")
GREY = "666666"
COL_LINE, BLOCK_LINE = Side(style="thin", color="B4C6E7"), Side(style="thin", color="5B7DB1")


def lines(ws, first, last, cols):
    """Lines for a block of rows (a set, a combo's people): between the columns, and a darker one under the block."""
    for r in range(first, last + 1):
        for col in range(1, cols + 1):
            ws.cell(r, col).border = Border(left=COL_LINE, right=COL_LINE, bottom=BLOCK_LINE if r == last else None)


def header(ws, names, cols=None):
    """A table's column names, bold on blue, boxed."""
    ws.append(names)
    for c in ws[ws.max_row][:cols or len(names)]:
        c.font, c.fill = Font(name=FONT, bold=True), HEAD_FILL
        c.border = Border(left=COL_LINE, right=COL_LINE, top=BLOCK_LINE, bottom=BLOCK_LINE)


def open_label(settings):
    """What an open set says in the exports."""
    return "OPEN - volunteer" if settings.extra_slot_policy == "open" else "(empty)"


def write_schedule_xlsx(path, schedule, combos, settings, name_of, instruments):
    """schedule: a schedule_file.Schedule. combos: {id: Combo}; name_of(email) -> name; instruments:
    {(combo name, email): instrument} (Store.instruments). PermissionError: it's open in Excel."""
    wb = Workbook()
    wb.properties.subject = schedule.semester
    supervised = schedule.supervised or set()
    open_text = open_label(settings)

    def font(**kw):
        return Font(name=FONT, **kw)

    # By night: one table per night
    ws = wb.active
    ws.title = "By night"
    ws.append([f"{schedule.semester}: combo shows"])
    ws["A1"].font = font(bold=True, size=14)
    ws.append([f"Made by Combo Manager on {datetime.now():%Y-%m-%d %H:%M}. Editing this file changes nothing: "
               "make changes in the app, which rebuilds it."])
    ws["A2"].font = font(italic=True, color=GREY)
    for n in schedule.nights:
        ws.append([])
        fac = schedule.faculty.get(n.date)
        ws.append([f"{make_label(n.date)}  ·  {n.venue}" + (("  ·  FB" + (f": {fac.full}" if fac else ""))
                                                            if n.date in supervised else "")])
        r = ws.max_row
        for col in range(1, 6):
            ws.cell(r, col).fill = NIGHT_FILL
        ws.cell(r, 1).font = font(bold=True, size=11)
        header(ws, ["Set", "Time", "Combo", "Name", "Instrument"])
        for k in range(1, n.n_slots + 1):
            first = ws.max_row + 1
            cid, text = schedule.sets.get(n.date, {}).get(k), schedule.typed.get(n.date, {}).get(k)
            when = n.set_range(k)
            if cid:
                combo = combos[cid]
                inst = {e: instruments.get((combo.name, e), "") for e in combo.members}
                people = [(name_of(e) + ("  (liaison)" if e == combo.liaison else ""), inst[e])
                          for e in by_instrument(combo.members, inst.get, name_of)]
                if combo.professor:
                    people.append((name_of(combo.professor) + "  (coach)", ""))
                for i, (who, instrument) in enumerate(people or [("", "")]):
                    ws.append([k, when, combo.name] if i == 0 else ["", "", ""])
                    ws.cell(ws.max_row, 4).value, ws.cell(ws.max_row, 5).value = who, instrument
                    for c in ws[ws.max_row]:
                        c.font = font(bold=c.column == 3)
            else:
                ws.append([k, when, text or open_text])
                for c in ws[ws.max_row]:
                    c.font = font(italic=bool(text), color=None if text else GREY)
            lines(ws, first, ws.max_row, 5)
    for letter, w in zip("ABCDE", (6, 18, 16, 34, 16)):
        ws.column_dimensions[letter].width = w
    ws.page_setup.fitToWidth, ws.sheet_properties.pageSetUpPr.fitToPage = 1, True
    ws.page_setup.fitToHeight = 0

    def sheet(title, header, rows, widths):
        s = wb.create_sheet(title)
        s.append(header)
        for c in s[1]:
            c.font, c.fill = font(bold=True), HEAD_FILL
        for row in rows:
            s.append(row)
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.font = font()
        for letter, w in zip("ABCDEFGHIJ", widths):
            s.column_dimensions[letter].width = w
        s.freeze_panes = "A2"
        s.auto_filter.ref = s.dimensions
        return s

    timed = any(n.first_set is not None for n in schedule.nights)
    rows = []
    for n in schedule.nights:
        for k in range(1, n.n_slots + 1):
            cid, text = schedule.sets.get(n.date, {}).get(k), schedule.typed.get(n.date, {}).get(k)
            rows.append([n.date, WEEKDAYS[n.date.weekday()], n.venue, k]
                        + ([n.set_time(k), n.set_end(k)] if timed else [])
                        + [combos[cid].name if cid else text or open_text]
                        + (["Yes" if n.date in supervised else ""] if schedule.supervised is not None else []))
    s = sheet("All sets", ["Date", "Day", "Venue", "Set"] + (["Start", "End"] if timed else []) + ["Combo"]
              + (["FB"] if schedule.supervised is not None else []), rows,
              (14, 12, 14, 6) + ((10, 10) if timed else ()) + (36, 12))
    for r in range(2, s.max_row + 1):
        s.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        if timed:
            s.cell(row=r, column=5).number_format = s.cell(row=r, column=6).number_format = "h:mm AM/PM"
    if schedule.supervised is not None:
        sup = [[n.date, WEEKDAYS[n.date.weekday()], n.venue,
                *((schedule.faculty[n.date].full, schedule.faculty[n.date].email) if n.date in schedule.faculty
                  else ("", "")),
                ", ".join(combos[c].name for k, c in sorted(schedule.sets.get(n.date, {}).items()) if c)]
               for n in schedule.nights if n.date in supervised]
        s = sheet("Feedback nights", ["Date", "Day", "Venue", "Faculty member", "Faculty email", "Combos playing"],
                  sup, (14, 12, 14, 22, 30, 60))
        for r in range(2, s.max_row + 1):
            s.cell(row=r, column=1).number_format = "yyyy-mm-dd"
    rows = []
    for entry in reversed(schedule.history):
        saved = entry.get("saved", "").replace("T", " ")[:16]
        what = "; ".join(entry.get("what", []))
        cells = entry.get("changes") or [{}]
        for i, ch in enumerate(cells):
            night = ch.get("night")
            rows.append([saved if i == 0 else "", entry.get("computer", "") if i == 0 else "", what if i == 0 else "",
                         make_label(datetime.fromisoformat(night).date()) if night else "", ch.get("set", ""),
                         ch.get("before", ""), ch.get("after", "")])
    if rows:
        sheet("Changes", ["Saved", "Computer", "What", "Night", "Set", "Before", "After"], rows,
              (17, 14, 70, 14, 6, 20, 20))
    if schedule.report:
        sheet("Report", ["Level", "Message"], [[str(l).upper(), t] for l, t in schedule.report], (10, 120))
    save_workbook(wb, path)


# ---------------------------------------------------------------- the old, hand-editable Schedule.xlsx (read once)

def is_old_schedule(path):
    """True for a Schedule.xlsx of the old kind (it has a 'Schedule' sheet; the exports don't)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path, read_only=True)
    except Exception:                                  # not a workbook at all, or locked: not ours to convert
        return False
    try:
        return "Schedule" in wb.sheetnames
    finally:
        wb.close()


def old_schedule_semester(path, settings=None):
    """The semester an old Schedule.xlsx was made for (kept in the file's properties). For a file made before that was
    recorded: settings.semester_name if all its nights fall within the settings' dates, else None (unknown)."""
    from openpyxl import load_workbook
    from util import to_date
    try:
        wb = load_workbook(path, read_only=True)
    except FileNotFoundError:
        return None
    try:
        if wb.properties.subject:
            return wb.properties.subject
        if settings is None or "Schedule" not in wb.sheetnames:
            return None
        dates = []
        for r in wb["Schedule"].iter_rows(min_row=2, max_col=1, values_only=True):
            try:
                if r[0] is not None:
                    dates.append(to_date(r[0]))
            except ValueError:
                pass
        inside = dates and all(settings.start_date <= d <= settings.end_date for d in dates if d)
        return settings.semester_name if inside else None
    finally:
        wb.close()


def read_old_schedule(path, combos):
    """An old Schedule.xlsx (possibly edited by hand) -> (sets, problems, supervised, typed).
    sets = {date: {set number: combo id or None}}. A Combo cell that is blank, 'OPEN...' or '(empty)' is open (None).
    Anything else that isn't a combo name (e.g. 'Jam session') goes in typed = {date: {set number: text}}: shown on
    the PDF as written, and the set counts as taken (None in sets). Text that looks like a mistyped combo name
    ('Combo 5') is reported instead. supervised = the dates with Yes in the FB column, or None if the file
    has no such column."""
    from openpyxl import load_workbook
    from util import blank, to_date
    try:
        wb = load_workbook(path, data_only=True)
    except FileNotFoundError:
        raise ScheduleFileError(f"No {Path(path).name} yet: make the schedule first (Run tab, step 2).")
    if "Schedule" not in wb.sheetnames:
        raise ScheduleFileError(f"{path} has no 'Schedule' sheet.")
    ws = wb["Schedule"]
    by_name = {c.name.strip().casefold(): c.id for c in combos.values()}
    head = [str(v or "").strip().lower() for v in next(ws.iter_rows(max_row=1, values_only=True), ())]
    di, si, ci = (head.index(h) if h in head else j for h, j in (("date", 0), ("set", 3), ("combo", 4)))
    vi = next((head.index(h) for h in ("fb", "supervised") if h in head), None)   # (older files: Supervised)
    sets, problems, supervised, typed = defaultdict(dict), [], set(), defaultdict(dict)
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        r = tuple(r) + (None,) * (max(di, si, ci, vi or 0) + 1 - len(r))
        if all(blank(v) for v in (r[di], r[si], r[ci])):
            continue
        try:
            d, k = to_date(r[di]), int(r[si])
        except (ValueError, TypeError):
            problems.append(f"Schedule row {i}: can't read the date or set number; row ignored.")
            continue
        if vi is not None and str(r[vi] or "").strip().lower() in ("yes", "y", "x", "true"):
            supervised.add(d)
        name = "" if r[ci] is None else str(r[ci]).strip()
        if not name or name.upper().startswith("OPEN") or name == "(empty)":
            cid = None
        elif name.casefold() in by_name:
            cid = by_name[name.casefold()]
        elif re.match(r"(?i)combo\s*\d", name):
            problems.append(f"Schedule row {i}: '{name}' isn't an accepted combo (withdrawn, or a typo?). Treated "
                            "as open: give the set away or leave it open (Schedule or Swaps tab).")
            cid = None
        else:
            typed[d][k] = name
            cid = None
        if k in sets[d]:
            problems.append(f"Schedule row {i}: {d} set {k} appears twice; using the later row.")
        sets[d][k] = cid
    return dict(sets), problems, (supervised if vi is not None else None), dict(typed)


def old_schedule_nights(path, settings):
    """The show nights as an old Schedule.xlsx has them (date, venue, number of sets, set times), sorted."""
    from datetime import datetime as dt
    from openpyxl import load_workbook
    from core.model import Night
    from util import to_date, to_time
    try:
        ws = load_workbook(path, data_only=True)["Schedule"]
    except (FileNotFoundError, KeyError):
        raise ScheduleFileError(f"Can't read the nights from {path}.")
    head = [str(v or "").strip().lower() for v in next(ws.iter_rows(max_row=1, values_only=True), ())]
    col = {h: head.index(h) for h in ("date", "venue", "set", "start", "end") if h in head}
    rows = defaultdict(dict)                          # date -> {set: (venue, start, end)}
    for r in ws.iter_rows(min_row=2, values_only=True):
        r = tuple(r) + (None,) * (len(head) + 1 - len(r))
        try:
            d, k = to_date(r[col.get("date", 0)]), int(r[col.get("set", 3)])
            start = to_time(r[col["start"]]) if "start" in col else None
            end = to_time(r[col["end"]]) if "end" in col else None
        except (ValueError, TypeError):
            continue                                  # read_old_schedule reports unreadable rows
        if d:
            rows[d][k] = (str(r[col.get("venue", 2)] or "").strip(), start, end)
    # as in core.slots: a regular show day's own flag; any other night is preferred when its venue is
    by_day = {(sd.weekday, sd.venue.casefold()): sd.supervision_preferred for sd in settings.show_days}
    preferred_venues = {sd.venue.casefold() for sd in settings.show_days if sd.supervision_preferred}

    def minutes(a, b):
        return int((dt.combine(dt.min, b) - dt.combine(dt.min, a)).total_seconds() // 60)
    nights = []
    for d, sets in sorted(rows.items()):
        venue, first, end1 = sets[min(sets)]
        length = minutes(first, end1) if first and end1 else 0
        gap = minutes(end1, sets[2][1]) if first and end1 and 2 in sets and sets[2][1] else 0
        nights.append(Night(d, venue, max(sets), first if length > 0 else None, max(length, 0), max(gap, 0),
                            by_day.get((d.weekday(), venue.casefold()), venue.casefold() in preferred_venues)))
    return nights
