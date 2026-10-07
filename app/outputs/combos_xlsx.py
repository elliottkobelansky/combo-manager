"""Writes Combos.xlsx (the Combos tab's Export Excel): the combo list, looking like the Schedule.xlsx export. One
table per combo (first-year combos and shows in its heading): each member in the Combos tab's order (by instrument;
the liaison marked) with instrument and email, then the supervisor.

An export: never read back, so editing it changes nothing.
"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Border, Font

from outputs.excel_schedule import BLOCK_LINE, COL_LINE, FONT, GREY, NIGHT_FILL, header
from shared_folder import save_workbook
from util import by_instrument


def write_combos_xlsx(path, combos, name_of, semester, instruments=None, shows=None):
    """combos in the order to list; name_of(email) -> name; instruments = {(combo name, email): instrument};
    shows = {combo id: ['Tue Oct 13', 'Fri Nov 6*', ...]} (* = supervised night). PermissionError: open in Excel."""
    instruments, shows = instruments or {}, shows or {}

    def font(**kw):
        return Font(name=FONT, **kw)

    def people(c):
        """(name, instrument, email, kind) per member ('liaison' or ''), then the supervisor ('supervisor')."""
        rows = [(name_of(e) + ("  (liaison)" if e == c.liaison else ""), instruments.get((c.name, e), ""), e,
                 "liaison" if e == c.liaison else "")
                for e in by_instrument(c.members, lambda e: instruments.get((c.name, e), ""), name_of)]
        if c.professor:
            rows.append((name_of(c.professor) + "  (coach)", "", c.professor, "supervisor"))
        return rows

    wb = Workbook()
    wb.properties.subject = semester
    ws = wb.active
    ws.title = "Combos"
    ws.append([f"{semester}: combos"])
    ws["A1"].font = font(bold=True, size=14)
    ws.append([f"Exported from Combo Manager on {datetime.now():%Y-%m-%d %H:%M}. Editing this file changes "
               "nothing: make changes in the app (Combos tab), then export again."])
    ws["A2"].font = font(italic=True, color=GREY)
    for c in combos:
        ws.append([])
        dates = [d.replace("*", " (feedback night)") for d in shows.get(c.id, [])]
        ws.append([c.name + ("  ·  first-year combo" if c.first_year else "")
                   + (f"  ·  shows: {', '.join(dates)}" if dates else "")])
        for col in range(1, 4):
            ws.cell(ws.max_row, col).fill = NIGHT_FILL
        ws.cell(ws.max_row, 1).font = font(bold=True, size=11)
        header(ws, ["Name", "Instrument", "Email"])
        rows = people(c)
        for i, (name, instrument, email, kind) in enumerate(rows):
            ws.append([name, instrument, email])
            last = i == len(rows) - 1
            for cell in ws[ws.max_row]:
                cell.font = font(italic=kind == "supervisor")
                cell.border = Border(left=COL_LINE, right=COL_LINE, bottom=BLOCK_LINE if last else COL_LINE)
        if not c.professor:
            ws.append(["No coach yet"])
            ws.cell(ws.max_row, 1).font = font(italic=True, color="C00000")
    for letter, w in zip("ABC", (34, 16, 34)):
        ws.column_dimensions[letter].width = w
    ws.page_setup.fitToWidth, ws.sheet_properties.pageSetUpPr.fitToPage = 1, True
    ws.page_setup.fitToHeight = 0
    save_workbook(wb, path)
