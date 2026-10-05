"""Output writer: core.Result -> Schedule.xlsx, and reading it back after hand edits."""
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill



def write_schedule(path, result, settings):
    wb = Workbook()
    hfont, hfill, base = Font(name="Arial", bold=True), PatternFill("solid", fgColor="DDEBF7"), Font(name="Arial")
    first = [True]

    def sheet(title, header, rows, widths):
        ws = wb.active if first[0] else wb.create_sheet()
        first[0] = False
        ws.title = title
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = hfont, hfill
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = base
        for letter, w in zip("ABCDEFGHIJ", widths):
            ws.column_dimensions[letter].width = w
        ws.freeze_panes = "A2"
        return ws

    policy, combos = settings.extra_slot_policy, result.combos

    timed = any(n.first_set is not None for n in result.nights)   # Start/End columns only when set times are set up
    rows = []
    for n in result.nights:
        order = result.lineup.get(n.date, [])
        for k in range(1, n.n_slots + 1):
            c = order[k - 1] if k <= len(order) else None
            rows.append([n.date, n.weekday, n.venue, k] + ([n.set_time(k), n.set_end(k)] if timed else [])
                        + [combos[c].name if c else ("OPEN - volunteer" if policy == "open" else "(empty)")]
                        + (["Yes" if n.date in result.supervised else ""] if settings.every_combo_supervised else []))
    header = (["Date", "Day", "Venue", "Set"] + (["Start", "End"] if timed else []) + ["Combo"]
              + (["Supervised"] if settings.every_combo_supervised else []))
    ws = sheet("Schedule", header, rows, (14, 12, 14, 6) + ((10, 10) if timed else ()) + (36, 12))
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        if timed:
            ws.cell(row=r, column=5).number_format = ws.cell(row=r, column=6).number_format = "h:mm AM/PM"

    if settings.every_combo_supervised:
        sheet("Supervision", ["Date", "Day", "Venue", "Professor (fill in)"],
              [[n.date, n.weekday, n.venue, ""] for n in result.nights if n.date in result.supervised],
              (14, 12, 14, 30))
        for r in range(2, wb["Supervision"].max_row + 1):
            wb["Supervision"].cell(row=r, column=1).number_format = "yyyy-mm-dd"

    sheet("Report", ["Level", "Message"], [[l.upper(), t] for l, t in result.report], (10, 120))
    wb.save(path)


def read_schedule(path, combos):
    """Schedule.xlsx (possibly edited by hand) -> (sets, problems, supervised, typed).
    sets = {date: {set number: combo id or None}}. A Combo cell that is blank, 'OPEN...' or '(empty)' is open (None).
    Anything else that isn't a combo name (e.g. 'Jam session') goes in typed = {date: {set number: text}}: shown on
    the PDF as written, and the set counts as taken (None in sets). Text that looks like a mistyped combo name
    ('Combo 5') is reported instead. supervised = the dates with Yes in the Supervised column, or None if the file
    has no such column."""
    from openpyxl import load_workbook
    from util import blank, to_date
    try:
        wb = load_workbook(path, data_only=True)
    except FileNotFoundError:
        raise ScheduleFileError(f"Can't find {path}. Run solve.py first.")
    if "Schedule" not in wb.sheetnames:
        raise ScheduleFileError(f"{path} has no 'Schedule' sheet.")
    ws = wb["Schedule"]
    by_name = {c.name.strip().casefold(): c.id for c in combos.values()}
    head = [str(v or "").strip().lower() for v in next(ws.iter_rows(max_row=1, values_only=True), ())]
    di, si, ci = (head.index(h) if h in head else j for h, j in (("date", 0), ("set", 3), ("combo", 4)))
    vi = head.index("supervised") if "supervised" in head else None
    sets, problems, supervised, typed = defaultdict(dict), [], set(), defaultdict(dict)
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        r = tuple(r) + (None,) * (max(di, si, ci, vi or 0) + 1 - len(r))
        if all(blank(v) for v in (r[di], r[si], r[ci])):
            continue
        try:
            d, k = to_date(r[di]), int(r[si])
        except (ValueError, TypeError):
            problems.append(f"Schedule row {i}: can't read the date or set number; row ignored.")
            continue
        if vi is not None and str(r[vi] or "").strip().lower() in ("yes", "y", "x", "true"):
            supervised.add(d)
        name = "" if r[ci] is None else str(r[ci]).strip()
        if not name or name.upper().startswith("OPEN") or name == "(empty)":
            cid = None
        elif name.casefold() in by_name:
            cid = by_name[name.casefold()]
        elif re.match(r"(?i)combo\s*\d", name):
            problems.append(f"Schedule row {i}: '{name}' looks like a combo but isn't an accepted one (a typo?). "
                            "Treated as open.")
            cid = None
        else:
            typed[d][k] = name
            cid = None
        if k in sets[d]:
            problems.append(f"Schedule row {i}: {d} set {k} appears twice; using the later row.")
        sets[d][k] = cid
    return dict(sets), problems, (supervised if vi is not None else None), dict(typed)


class ScheduleFileError(Exception):
    pass


def write_swap(path, changes, combos, open_label):
    """Writes swapped sets into the Combo column of Schedule.xlsx (everything else, including hand edits, stays).
    changes = {(date, set number): combo id or None (open)}. A copy of the file as it was goes into a
    'Schedule backups' folder next to it first. Returns the backup's path."""
    import shutil
    from datetime import datetime
    from openpyxl import load_workbook
    from util import to_date
    path = Path(path)
    backups = path.parent / "Schedule backups"
    backups.mkdir(exist_ok=True)
    backup = backups / f"{path.stem} {datetime.now():%Y-%m-%d %H%M%S}{path.suffix}"
    shutil.copy2(path, backup)
    wb = load_workbook(path)
    ws = wb["Schedule"]
    head = [str(v or "").strip().lower() for v in next(ws.iter_rows(max_row=1, values_only=True), ())]
    di, si, ci = (head.index(h) if h in head else j for h, j in (("date", 0), ("set", 3), ("combo", 4)))
    todo = dict(changes)
    for row in ws.iter_rows(min_row=2):
        try:
            key = (to_date(row[di].value), int(row[si].value))
        except (ValueError, TypeError):
            continue
        if key in todo:
            c = todo.pop(key)
            row[ci].value = combos[c].name if c else open_label
    if todo:
        raise ScheduleFileError("Couldn't find these sets in Schedule.xlsx: "
                                + ", ".join(f"{d} set {k}" for d, k in sorted(todo)) + ". Nothing was changed.")
    try:
        wb.save(path)
    except PermissionError:
        raise ScheduleFileError(f"Can't save {path.name}: it's open in Excel. Close it and try again. "
                                "Nothing was changed.")
    return backup
