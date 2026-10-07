"""Makes a fake approvals table (Approvals.xlsx, in the layout the Power Automate flow writes) and a fake
conflicts table (Conflicts.xlsx, in the layout the conflict flow writes), with deliberate mistakes, for testing.
make() also returns what the approvals should turn into (worked out from what was put in, not by parsing), so
test_rules.py can check inputs.py against it.

    python dev/make_fake_forms.py FOLDER [--combos 33] [--seed 1]  # a demo data folder: fake spreadsheets, instruments
    python dev/make_fake_forms.py FOLDER --instruments [--replace]  # only fill in instruments for the combos there
FOLDER is any folder (made if needed), e.g. ~/Documents/Combo Manager demo; choose it in the app with Change
folder... to try things. Never inside the program folder: that holds no data. Without settings yet, it gets the
example settings (settings_file.DEFAULTS). Instruments go into AppFiles/scheduler_data.json (as if set in the Combos
tab); people who already have one keep it unless --replace.
"""
import argparse
import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "app"))  # the scheduler's code

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from core import generate_nights
from data_folder import APPROVALS_FILE, CONFLICTS_FILE, settings_path, store_path
from settings_file import DEFAULTS, load_settings, save_data

EPOCH = date(1899, 12, 30)
OUTSIDE = ["jamie.outside@gmail.com", "sam.guest@gmail.com"]   # non-McGill members: must be kept like anyone else
# Made-up faculty (not real people), staff addresses (@mcgill.ca); picked by Response Id, so the random draws for
# members and conflicts stay the same as before.
PROFS = ["helene.marchand", "darnell.whitfield", "ingrid.solberg", "rafael.quintero", "catherine.beaulieu",
         "malcolm.ashby", "yuki.tanabe", "olivier.gauthier"]


def prof_for(rid):
    return f"{PROFS[rid % len(PROFS)]}@mcgill.ca"

FIRST = ["Ana", "Ben", "Chloe", "Daniel", "Emma", "Felix", "Grace", "Hugo", "Isla", "Jonah", "Kenji", "Leah", "Mateo",
         "Nadia", "Oscar", "Priya", "Quinn", "Rosa", "Samir", "Tess", "Uma", "Victor", "Wen", "Xavier", "Yara", "Zoe",
         "Amir", "Bea", "Camille", "Dev", "Elodie", "Farah", "Gabriel", "Hana", "Ines", "Jules", "Kai", "Lucia",
         "Marc", "Noor", "Olivia", "Pablo", "Rhea", "Simon", "Talia", "Theo", "Vera", "Will", "Yusuf", "Zara",
         "Eloise", "Jean-Luc", "Mei", "Oren", "Sofia", "Luca", "Maya", "Niko", "Aditi", "Rafael"]
LAST = ["Ruiz", "Chen", "Tremblay", "Nguyen", "Gagnon", "Patel", "Roy", "Kim", "Cote", "Singh", "Bouchard", "Haddad",
        "Lee", "Morin", "Okafor", "Fortin", "Garcia", "Levesque", "Cohen", "Wong", "Pelletier", "Ali", "Rossi", "Dubois",
        "Silva", "Lavoie", "Kowalski", "Bergeron", "Mendes", "Ito", "Ferreira", "O'Brien", "Novak", "Gauthier",
        "Khan", "Fischer", "Leblanc", "Moreau", "Santos", "Park", "Hebert", "Murphy", "Costa", "Jensen", "Cyr"]


def fake_emails(n):
    """n McGill-style student emails (first.last@mail.mcgill.ca) for distinct, realistic names. The same list every
    time, and a longer list starts with a shorter one. Index 3 and 4 are two different people both called Alex
    Martin (alex.martin@ and alex.martin2@), as happens in real life."""
    rnd, seen, out = random.Random(2027), set(), []
    while len(out) < n:
        if len(out) == 3:
            out += ["alex.martin@mail.mcgill.ca", "alex.martin2@mail.mcgill.ca"]
            continue
        first, last = rnd.choice(FIRST), rnd.choice(LAST)
        key = f"{first}.{last}".lower().replace("'", "")
        if key in seen or key == "alex.martin":
            continue
        seen.add(key)
        out.append(f"{key}@mail.mcgill.ca")
    return out[:n]


def stamp(i):
    return f"10/{3 + i // 600}/2026 {16 + (i // 60) % 8}:{i % 60:02d}"


def date_cell(rnd, d):
    """The same date in the three shapes it can show up as: text (like the export), real date, Excel number."""
    r = rnd.random()
    if r < 0.7:
        return f"{d.month}/{d.day}/{d.year}"
    if r < 0.9:
        return datetime(d.year, d.month, d.day)
    return (d - EPOCH).days


def clean(e):
    e = e.lower()
    return e[: -len("@mcgill.ca")] + "@mail.mcgill.ca" if e.endswith("@mcgill.ca") else e


def expected_combos(truth, sem):
    """What the accepted combos should be: accepted rows of this semester, identical members -> the later row, in
    Response Id order. [(response id, members, liaison, first year)]"""
    latest = {}
    for t in truth:
        if t["semester"] == sem and t["status"].lower() == "accepted":
            latest[frozenset(t["members"])] = t
    return [(str(t["rid"]), frozenset(t["members"]), t["members"][0], t["fy"] == "Yes")
            for t in sorted(latest.values(), key=lambda t: t["rid"])]


def make(approvals_path, conflicts_path, nights, n_combos=33, seed=1, sem="Fall 2026", mistakes=True):
    rnd = random.Random(seed)
    dates = [n.date for n in nights]
    days = [dates[0] + timedelta(days=i) for i in range((dates[-1] - dates[0]).days + 1)]
    n_students = int(n_combos * 5 / 1.3)                      # ~30% of students play in two combos
    emails = fake_emails(n_students)
    pool = emails[:]
    rnd.shuffle(pool)

    # ---------------- approvals table (written by the flow, decided by the director) ----------------
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Response Id", "Submitted", "Semester", "Liaison", "Members", "Supervisor", "First year", "Status",
               "Decided by", "Decided on", "Notes"])
    groups = [[] for _ in range(n_combos)]
    for k, e in enumerate(pool):
        groups[k % n_combos].append(e)
    for e in rnd.sample(emails, n_combos * 5 - n_students):  # second memberships
        rnd.choice(groups).append(e)
    rid, truth = 0, []

    def add(members, semester=sem, fy="No", status="Accepted", note=None):
        nonlocal rid
        rid += 1
        members = list(dict.fromkeys(members))
        truth.append(dict(rid=rid, semester=semester, fy=fy, status=status,
                          members=list(dict.fromkeys(clean(e) for e in members))))
        decided = status not in ("Pending", "")
        ws.append([rid, f"2026-10-{3 + rid // 40:02d}T{16 + rid % 8}:{rid % 60:02d}:00Z", semester, members[0],
                   "\n".join(members), prof_for(rid), fy, status, "Director" if decided else None,
                   f"2026-10-{4 + rid // 40:02d}" if decided else None, note])

    if mistakes:
        for k in range(3):                                                  # last semester's rows come first
            add([f"old{k}@mail.mcgill.ca", f"old{k + 10}@mail.mcgill.ca"], semester="Winter 2026")
        add(["test@mail.mcgill.ca", "test2@mail.mcgill.ca"], semester=None, status="Pending")   # test row
    for gi, g in enumerate(groups):
        g = list(g)
        if mistakes and gi == 3:
            g[1] = g[1].replace("@mail.mcgill.ca", "@MCGILL.ca")            # sloppy domain: gets fixed
        if mistakes and gi == 5:
            g.append(OUTSIDE[0])                                            # Gmail member
        if mistakes and gi == 6:
            g.insert(0, OUTSIDE[1])                                         # Gmail liaison
        status = "accepted" if mistakes and gi % 5 == 4 else "Accepted"     # capitalisation doesn't matter
        add(g, fy="Yes" if rnd.random() < 0.25 else "No", status=status)
    if mistakes:
        add(groups[0])                                                      # accepted twice: latest wins
        add(groups[1][:-1] + ["late.joiner@mail.mcgill.ca"], status="Rejected", note="resubmitted")
        add(rnd.sample(emails, 4), status="Pending")                        # no decision yet
        add(rnd.sample(emails, 3), status="Withdrawn", note="left the program")
        add(rnd.sample(emails, 3), status="Maybe")                          # not a valid status
        add(rnd.sample(emails, 3), status="Approve")                        # raw approval outcome: not valid either
        add(["test3@mail.mcgill.ca"], semester=None, status="Pending")      # newest row has no semester: the
                                                                            # semester guess must skip it
    ws.add_table(Table(displayName="Approvals", ref=f"A1:K{ws.max_row}"))
    wb.save(approvals_path)

    # ---------------- conflicts table (written by the conflict flow; Status edited by the director) ----------------
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Response ID", "Submitted", "Semester", "Email", "Date 1", "Date 2", "Date 3", "Date 4", "Reason",
               "Additional Info", "Status", "Notes"])
    rid = 0

    def addc(email, ds, semester=sem, status="Active", note=None):
        nonlocal rid
        rid += 1
        cells = [date_cell(rnd, d) if isinstance(d, date) else d for d in ds[:4]]
        ws.append([rid, stamp(rid), semester, email] + cells + [None] * (4 - len(cells))
                  + ["exam" if rid % 3 else None, None, status, note])

    for e in emails:
        k = rnd.choices([0, 1, 2, 3, 4], weights=[38, 25, 17, 10, 10])[0]
        if k == 0:
            continue                                                         # no conflicts -> doesn't submit
        ds = [rnd.choice(dates) if rnd.random() < 0.8 else rnd.choice(days) for _ in range(k)]
        if mistakes and e == emails[7]:                                      # resubmission: first row gets replaced
            addc(e, [rnd.choice(dates) for _ in range(3)])
        addc(e if not (mistakes and e == emails[9]) else e.upper(), ds)
    if mistakes:
        addc(emails[1].replace("mail.", ""), [dates[2]])                    # bare @mcgill.ca
        addc(emails[2], [dates[3], "not a date", dates[4]])                 # garbage cell
        addc("nobody@mail.mcgill.ca", [dates[0]])                           # in no combo
        addc("Jamie.Outside@Gmail.com", [dates[5]])                         # Gmail member, typed in mixed case
        addc(OUTSIDE[1], [dates[8]])                                        # Gmail liaison
        addc(emails[11], dates[:4], status="Overruled", note="no reason given")   # director overruled it
        addc(emails[12], [dates[9]], status="ok")                           # synonym for Active
        addc("", [dates[1]])                                                # no email -> skipped
        addc(emails[5], [dates[6]], semester="Winter 2026")                 # old semester
        addc(emails[6], [dates[6]], semester=None)                          # no semester: test row
    ws.add_table(Table(displayName="Conflicts", ref=f"A1:L{ws.max_row}"))
    wb.save(conflicts_path)
    return expected_combos(truth, sem)


HORNS = [("Saxophone", 40), ("Trumpet", 28), ("Trombone", 14), ("Voice", 12), ("other", 6)]
OTHER = ["Vibraphone", "Flute", "Violin", "Clarinet"]


def lineup(n, rnd):
    """A realistic instrumentation for an n-piece combo: a rhythm section, then horns (and the odd singer)."""
    if n <= 2:
        return rnd.choice([["Piano", "Bass"], ["Guitar", "Voice"], ["Piano", "Voice"]])[:n]
    roles = ["Drums", "Bass"]
    roles += ["Piano"] if rnd.random() < 0.7 else ["Guitar"]
    if n >= 5 and rnd.random() < 0.3:
        roles.append("Guitar" if roles[-1] == "Piano" else "Piano")
    while len(roles) < n:
        horn = rnd.choices([h for h, _ in HORNS], weights=[w for _, w in HORNS])[0]
        roles.append(rnd.choice(OTHER) if horn == "other" else horn)
    return roles


def fake_instruments(folder, settings, seed=1, replace=False):
    """Gives every member of the accepted combos in `folder` a realistic instrument (scheduler_data.json). Someone in
    several combos plays the same instrument in each when that combo needs it (otherwise they double on another).
    Returns how many were set."""
    from inputs import load_input
    from store import Store
    rnd = random.Random(seed)
    combos = sorted(load_input(folder, settings).combos, key=lambda c: c.name)
    store = Store(folder)
    sem = settings.semester_name
    have = {} if replace else store.instruments(sem)
    plays = {}                                     # email -> instrument, so people stay on one instrument
    for (_, e), inst in have.items():
        plays.setdefault(e, inst)
    n_set = 0
    for c in combos:
        roles, members = lineup(len(c.members), rnd), sorted(c.members)
        rnd.shuffle(members)
        given = {e: have[(c.name, e)] for e in members if (c.name, e) in have}       # 1. already set: kept
        for inst in given.values():
            if inst in roles:
                roles.remove(inst)
        for e in members:                    # 2. same instrument as in their other combos, if this one needs it
            if e not in given and plays.get(e) in roles:
                given[e] = plays[e]
                roles.remove(plays[e])
        for e in members:                                                            # 3. the roles still open
            if e not in given:
                given[e] = roles.pop(0) if roles else rnd.choice(["Saxophone", "Trumpet", "Voice"])
        for e, inst in given.items():
            plays.setdefault(e, inst)
            if (c.name, e) not in have:
                store.set_instrument(sem, c.name, e, inst)
                n_set += 1
    store.save()
    return n_set


def make_demo(folder, combos=33, seed=1, force=False):
    """A data folder to try the app with: the example settings (unless it has settings already), fake approvals and
    conflicts for that semester, and instruments. -> how many instruments were set."""
    folder = Path(folder)
    approvals, conflicts = folder / APPROVALS_FILE, folder / CONFLICTS_FILE
    existing = [f.name for f in (approvals, conflicts) if f.exists()]
    if existing and not force:
        raise SystemExit(f"{', '.join(existing)} already exist in {folder} (real data?). Use another folder, "
                         "or run with --force to overwrite them with fake data.")
    folder.mkdir(parents=True, exist_ok=True)
    if not settings_path(folder).exists():
        save_data(settings_path(folder), DEFAULTS)
    settings, _ = load_settings(settings_path(folder))
    make(approvals, conflicts, generate_nights(settings), combos, seed, settings.semester_name)
    return fake_instruments(folder, settings, seed)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("folder", type=Path, help="the demo data folder, e.g. ~/Documents/Combo Manager demo")
    ap.add_argument("--combos", type=int, default=33)
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--force", action="store_true", help="overwrite existing response files")
    ap.add_argument("--instruments", action="store_true", help="only fill in instruments for the combos there")
    ap.add_argument("--replace", action="store_true", help="with --instruments: also replace instruments already set")
    a = ap.parse_args()
    folder = a.folder.expanduser()
    if a.instruments:
        settings, _ = load_settings(settings_path(folder))
        n = fake_instruments(folder, settings, a.seed, a.replace)
        print(f"Set {n} instrument(s) in {store_path(folder)}.")
        sys.exit(0)
    n = make_demo(folder, a.combos, a.seed, a.force)
    print(f"Wrote {APPROVALS_FILE} and {CONFLICTS_FILE} in {folder}, and {n} instrument(s). In the app: Change "
          "folder..., pick it, then Make schedule.")
