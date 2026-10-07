"""Runs fake downloads -> inputs -> solve on fake data and checks every hard rule independently of the solver.

    python dev/test_rules.py
"""
import re
import sys
import tempfile
from collections import defaultdict
from dataclasses import replace
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "app"))  # the scheduler's code

from core import ScheduleError, generate_nights, run_schedule
from make_fake_forms import OUTSIDE, fake_emails, make
from openpyxl import load_workbook

from inputs import load_input
from settings_file import DEFAULTS, validate
from data_folder import app_data, lock_path


def check(result, inp, settings):
    bad, nights, combos = [], {n.date: n for n in result.nights}, result.combos
    show_min = defaultdict(int)
    for sd in settings.show_days:
        show_min[sd.venue] = max(show_min[sd.venue], sd.min_per_combo)
    cnt = defaultdict(lambda: defaultdict(int))
    for d, cs in result.lineup.items():
        n = nights[d]
        if len(cs) > n.n_slots:
            bad.append(f"{d}: more combos than sets")
        if len(cs) != len(set(cs)):
            bad.append(f"{d}: a combo plays twice in one night")
        for c in cs:
            cnt[c][n.venue] += 1
            for m in combos[c].members:
                if d in inp.blocked.get(m, ()):
                    bad.append(f"{combos[c].name} plays {d} but {m} is blocked")
    for c in combos:
        for v, m in show_min.items():
            if cnt[c][v] < m:
                bad.append(f"{combos[c].name} has {cnt[c][v]} {v} show(s), needs {m}")
        if settings.max_shows_per_combo and sum(cnt[c].values()) > settings.max_shows_per_combo:
            bad.append(f"{combos[c].name} exceeds the show cap")
        total, need = sum(cnt[c].values()), settings.min_shows_per_combo
        if need and total < need:
            bad.append(f"{combos[c].name} has {total} shows, needs {need}")
        if need and (settings.extra_slot_policy == "open" or combos[c].first_year):
            if total > max(need, sum(show_min.values())):
                bad.append(f"{combos[c].name} got extra shows it shouldn't (exactly {need} in open mode and for "
                           "first-year combos)")
    if settings.every_combo_supervised:
        for c in combos:
            if not any(c in cs for d, cs in result.lineup.items() if d in result.supervised):
                bad.append(f"{combos[c].name} plays no supervised night")
        for d in result.supervised:
            if len(result.lineup.get(d, [])) < nights[d].n_slots:
                bad.append(f"supervised night {d} has open sets")
        cap = settings.max_supervised_nights
        if cap and len(result.supervised) > cap:
            bad.append(f"{len(result.supervised)} supervised nights, cap is {cap}")
    return bad


def missed_venues(result, inp):
    """(combo name, venue) for each combo with 2+ shows and none at a venue where it had a usable night."""
    out = []
    for c, combo in sorted(result.combos.items()):
        blocked = set().union(*[inp.blocked.get(m, set()) for m in combo.members])
        mine = [n for n in result.nights if c in result.lineup.get(n.date, [])]
        if len(mine) < 2:
            continue
        for v in sorted({n.venue for n in result.nights}):
            usable = any(n.venue == v and n.date not in blocked for n in result.nights)
            if usable and not any(n.venue == v for n in mine):
                out.append((combo.name, v))
    return out


def check_approvals(combos, approvals_path):
    """Only Accepted rows may be scheduled, and every combo must carry its Response Id."""
    ws = load_workbook(approvals_path)["Sheet1"]
    status = {str(r[0]): str(r[7] or "").strip().lower() for r in ws.iter_rows(min_row=2, values_only=True)}
    return [f"{c.name} (response {c.ref or '?'}) is scheduled but its Status is '{status.get(c.ref)}'"
            for c in combos if status.get(c.ref) != "accepted"]


def check_outside(inp):
    """Non-McGill (Gmail) members stay in their combo, and their conflicts count."""
    members = set().union(*[c.members for c in inp.combos])
    return ([f"{e} was dropped from their combo" for e in OUTSIDE if e not in members]
            + [f"{e}'s conflicts were lost" for e in OUTSIDE if e not in inp.blocked])


def check_parse(inp, expected):
    """inputs.py reads the approvals table into exactly the combos make_fake_forms put in, numbered 01, 02, ..."""
    got = [(c.ref, c.members, c.liaison, c.first_year) for c in sorted(inp.combos, key=lambda c: c.name)]
    bad = [] if got == expected else [f"approvals read wrongly: {len(got)} combos, expected {len(expected)}"
                                      + next((f"; first difference: got {g}, expected {e}"
                                              for g, e in zip(got, expected) if g != e), "")]
    names = sorted(c.name for c in inp.combos)
    if names != [f"Combo {k:02d}" for k in range(1, len(names) + 1)]:
        bad.append(f"combo numbers aren't 01..{len(names):02d}: {names[:3]}...")
    return bad


def check_status(inp):
    """Overruled conflicts don't count; 'ok' (a synonym for Active) does. See make_fake_forms.py."""
    overruled, ok = fake_emails(13)[11], fake_emails(13)[12]
    return ([] if overruled not in inp.blocked else ["an Overruled conflict was counted"]) + \
           ([] if inp.blocked.get(ok) else ["a conflict with Status 'ok' was dropped"])


def main():
    tmp = Path(tempfile.mkdtemp())
    # (the fake sign-ups have @mcgill.ca slips, corrected by the domain fixes older settings may still have)
    base, _ = validate({**DEFAULTS, "email_domain_fixes": "mcgill.ca -> mail.mcgill.ca"})
    base = replace(base, solver_time_limit_sec=10)
    nights = generate_nights(base)
    failures = 0
    for n_combos, policy in [(8, "auto"), (22, "auto"), (22, "open"), (33, "auto"), (33, "open")]:
        s = replace(base, extra_slot_policy=policy)
        (app_data(tmp) / "scheduler_data.json").unlink(missing_ok=True)          # combo numbers start fresh per scenario
        expected = make(tmp / "a.xlsx", tmp / "f.xlsx", nights, n_combos, seed=1, sem=s.semester_name)
        inp = load_input(tmp, s, "a.xlsx", "f.xlsx")
        result = run_schedule(inp, s)
        bad = check(result, inp, s) + check_parse(inp, expected) + check_approvals(inp.combos, tmp / "a.xlsx") + check_outside(inp) + check_status(inp)
        one_venue = missed_venues(result, inp)
        if policy == "open":                    # one show per venue is always possible on this fake data
            bad += [f"{name} plays no {v} show but could have" for name, v in one_venue]
        counts = sorted(sum(1 for cs in result.lineup.values() if c in cs) for c in result.combos)
        print(f"{'PASS' if not bad else 'FAIL'}  {len(result.combos):>2} combos, {policy:<4} | "
              f"{result.stats['total_sets'] - result.stats['empty_sets']}/{result.stats['total_sets']} sets filled | "
              f"shows per combo {counts[0]}-{counts[-1]} | "
              f"soft rules broken: {len(result.stats['shared_nights'])} student double-night(s), "
              f"{len(result.stats['early_first_year'])} early first-year show(s) | "
              f"{len(result.supervised)} supervised nights | {len(one_venue)} combo(s) missing a venue they could play")
        for b in bad[:5]:
            print("      -", b)
        failures += bool(bad)
    # A combo that can't make any Clara night gets Upstairs twice instead (min_shows_per_combo), never Clara twice.
    s = replace(base, extra_slot_policy="open")
    who = inp.combos[0]
    clara = {n.date for n in nights if n.venue == "Clara"}
    blocked = dict(inp.blocked)
    blocked[sorted(who.members)[0]] = set(blocked.get(sorted(who.members)[0], set())) | clara
    result = run_schedule(replace(inp, blocked=blocked), s)
    mine = [n.venue for n in nights if who.id in result.lineup.get(n.date, [])]
    ok = sorted(mine) == ["Upstairs", "Upstairs"] and not check(result, replace(inp, blocked=blocked), s)
    print(f"{'PASS' if ok else 'FAIL'}  combo with no usable Clara night plays Upstairs twice ({', '.join(mine)})")
    failures += not ok

    # Swap finder: every option it offers must pass this file's own rule check; it must find some.
    from core.swaps import apply_option, swap_options
    s = replace(base, extra_slot_policy="open")
    result = run_schedule(inp, s)
    sets = {d: {k: c for k, c in enumerate(cs, start=1)} for d, cs in result.lineup.items()}
    tried = offered = overrides = 0
    bad = []
    samples = {}                                       # one option of each kind, for the summary check below
    breaking = []                                      # options that break a rule (for the summary check too)
    for c in sorted(result.combos)[:6]:
        for d, k in [(d, k) for d, row in sets.items() for k, x in row.items() if x == c][:1]:
            opts = swap_options(sets, result.nights, result.combos, inp, s, result.supervised, {}, c, d, k)
            if any(a.breaks and not b.breaks for a, b in zip(opts, opts[1:])):
                bad.append(f"{result.combos[c].name}: an option breaking a rule is listed before a legal one")
            for opt in opts:
                new = apply_option(sets, opt)
                lineup = {dd: [row[kk] for kk in sorted(row) if row[kk]] for dd, row in new.items()}
                found = check(replace(result, lineup=lineup), inp, s)
                if opt.breaks:                      # a manual override: it must really break something
                    overrides += 1
                    breaking.append((sets, result, opt, c))
                    if not found and not any("twice" in b or "supervis" in b for b in opt.breaks):
                        bad.append(f"{opt.title}: marked as breaking '{opt.breaks[0]}', but the check finds nothing")
                    continue
                samples.setdefault((opt.kind, opt.same_night), (sets, result, opt, c))
                bad += [f"{opt.title}: {b}" for b in found]
                offered += 1
            tried += 1
    ok = not bad and offered > 0 and overrides > 0
    print(f"{'PASS' if ok else 'FAIL'}  swap finder: {offered} options for {tried} shows keep the hard rules; "
          f"{overrides} more break one, listed last and named")
    for b in bad[:5]:
        print("      -", b)
    failures += not ok

    # Giving a show away: only a combo with more than it needs can, and every give-away keeps the hard rules.
    s = replace(base, extra_slot_policy="auto")
    result = run_schedule(inp, s)
    sets = {d: {k: c for k, c in enumerate(cs, start=1)} for d, cs in result.lineup.items()}
    count = {c: sum(1 for row in sets.values() for x in row.values() if x == c) for c in result.combos}
    rich = [c for c in sorted(result.combos) if count[c] > (s.min_shows_per_combo or 0)][:3]
    poor = [c for c in sorted(result.combos) if count[c] <= (s.min_shows_per_combo or 0)][:3]
    gives, bad, wrong = 0, [], []
    for c in rich + poor:
        for d, k in [(d, k) for d, row in sets.items() for k, x in row.items() if x == c]:
            for opt in swap_options(sets, result.nights, result.combos, inp, s, result.supervised, {}, c, d, k):
                if opt.kind not in ("give", "drop") or opt.breaks:
                    continue
                samples.setdefault((opt.kind, False), (sets, result, opt, c))
                if c in poor:
                    wrong.append(f"{result.combos[c].name} has only {count[c]} shows but may give one away")
                new = apply_option(sets, opt)
                lineup = {dd: [row[kk] for kk in sorted(row) if row[kk]] for dd, row in new.items()}
                for b in check(replace(result, lineup=lineup), inp, s):
                    # a first-year combo getting an extra show is a soft rule after publishing: must be warned about
                    if not ("got extra shows" in b and any("first-year" in w for w in opt.warnings)):
                        bad.append(f"{opt.title}: {b}")
                gives += 1
    # Claiming an open set (volunteering): every claim keeps the hard rules (the extra show itself is the point).
    s = replace(base, extra_slot_policy="open")
    result = run_schedule(inp, s)
    sets = {d: {k: c for k, c in enumerate(cs, start=1)} for d, cs in result.lineup.items()}
    claims = 0
    for c in sorted(result.combos)[:4]:
        for opt in swap_options(sets, result.nights, result.combos, inp, s, result.supervised, {}, c):
            if opt.breaks:
                continue
            samples.setdefault(("claim", False), (sets, result, opt, c))
            new = apply_option(sets, opt)
            lineup = {dd: [row[kk] for kk in sorted(row) if row[kk]] for dd, row in new.items()}
            bad += [f"{opt.title}: {b}" for b in check(replace(result, lineup=lineup), inp, s)
                    if "got extra shows" not in b]
            claims += 1
    ok = gives > 0 and claims > 0 and not bad and not wrong and rich and poor
    print(f"{'PASS' if ok else 'FAIL'}  give-aways: {gives} options from combos with extra shows, all legal; "
          f"none from combos at their minimum; {claims} open-set claims, all legal")
    for b in (bad + wrong)[:5]:
        print("      -", b)
    for b in bad[:5]:
        print("      -", b)
    failures += not ok

    # Swap summaries (the Swaps tab's "Copy swap summary"): both sides of every kind of option, in words.
    from core.swaps import involved, swap_summary
    headings = {("trade", False): "trade shows", ("trade", True): "swap set order", ("move", False): "moves to an open",
                ("move", True): "another set the same night", ("give", False): "gives a show to",
                ("drop", False): "gives up a show", ("claim", False): "takes an open set"}
    bad = []
    for key, (sets_, res, opt, c) in sorted(samples.items()):
        text = swap_summary(sets_, res.nights, res.combos, opt, c, lambda e: "Name of " + e.split("@")[0])
        ids = involved(sets_, opt, res.combos, c)
        if ids[0] != c or len(ids) != (2 if key[0] in ("trade", "give") else 1):
            bad.append(f"{key}: combos involved {ids}")
        if headings[key] not in text.splitlines()[0]:
            bad.append(f"{key}: heading '{text.splitlines()[0]}'")
        for x in ids:                                   # one line per combo touched: two-way
            if not any(line.startswith(f"\u2022 {res.combos[x].name} ") for line in text.splitlines()):
                bad.append(f"{key}: nothing about {res.combos[x].name}")
            if res.combos[x].liaison and "Name of " + res.combos[x].liaison.split("@")[0] not in text:
                bad.append(f"{key}: {res.combos[x].name}'s liaison missing")
        if "@" in text:
            bad.append(f"{key}: an email in the summary")
        if re.search(r"\b[a-z]+_[a-z_]+\b", text):
            bad.append(f"{key}: a setting's name in the summary")
    # a conflict overridden: the summary names the student, not the rule
    conflict = next(((sets_, res, o, c) for sets_, res, o, c in breaking if any("conflict" in b for b in o.breaks)),
                    None)
    if conflict:
        sets_, res, o, c = conflict
        text = swap_summary(sets_, res.nights, res.combos, o, c, lambda e: "Name of " + e.split("@")[0],
                            inp.blocked)
        if "had marked" not in text or "conflict" in text.split("Heads-up")[0] or "Breaks" in text:
            bad.append(f"a conflict override isn't explained for the combo: {text}")
    else:
        bad.append("no option breaking a conflict to check the summary with")
    missing = sorted(set(headings) - set(samples))
    ok = not bad and {("trade", False), ("give", False), ("drop", False), ("claim", False)} <= set(samples)
    print(f"{'PASS' if ok else 'FAIL'}  swap summaries: {len(samples)} kinds of option, both sides, liaison names, no emails or settings, a conflict named"
          + (f" (none found here: {', '.join(k for k, _ in missing)})" if missing else ""))
    for b_ in bad[:5]:
        print("      -", b_)
    failures += not ok

    # The schedule file (App data/schedule.json): saved and read back unchanged, changes with a backup first, text in
    # a set, a combo that's gone, an old hand-editable Schedule.xlsx converted, the xlsx export, archiving.
    import schedule_file as sf_
    from openpyxl import Workbook, load_workbook as open_wb
    from outputs.excel_schedule import is_old_schedule, write_schedule_xlsx
    bad = []
    sdir = tmp / "schedule test"
    sdir.mkdir()
    sf_.save_new(sdir, result, s)
    sch = sf_.load(sdir, result.combos, s)
    lineup = {d: [row[k] for k in sorted(row) if row[k]] for d, row in sch.sets.items()}
    if lineup != {d: cs for d, cs in result.lineup.items()} | {d: [] for d in lineup if d not in result.lineup}:
        bad.append("the saved schedule reads back different sets")
    if sch.supervised != set(result.supervised) or sch.semester != s.semester_name:
        bad.append("supervised nights or semester changed on the way")
    if [(n.date, n.venue, n.n_slots, n.first_set, n.set_length, n.break_minutes) for n in sch.nights] != \
            [(n.date, n.venue, n.n_slots, n.first_set, n.set_length, n.break_minutes) for n in result.nights]:
        bad.append("the nights changed on the way")
    (d1, k1), (d2, k2) = sorted((d, k) for d, row in sch.sets.items() for k, c in row.items() if c)[:2]
    c1, c2 = sch.sets[d1][k1], sch.sets[d2][k2]
    open_set = next(((d, k) for d, row in sch.sets.items() for k, c in row.items() if c is None), None)
    copy = sf_.save_changes(sdir, result.combos, sets={(d1, k1): c2, (d2, k2): c1})
    sch2 = sf_.load(sdir, result.combos, s)
    if sch2.sets[d1][k1] != c2 or sch2.sets[d2][k2] != c1 or not copy or not copy.exists():
        bad.append("a trade wasn't saved, or no backup was made first")
    if [h["what"] for h in sch2.history] != [["Schedule made"], []] or \
            sorted((ch["before"], ch["after"]) for ch in sch2.history[-1]["changes"]) != sorted(
                [(result.combos[c1].name, result.combos[c2].name), (result.combos[c2].name, result.combos[c1].name)]):
        bad.append(f"the history doesn't have the trade: {sch2.history}")
    if open_set:
        sf_.save_changes(sdir, result.combos, typed={open_set: "Jam session"})
        if sf_.load(sdir, result.combos, s).typed.get(open_set[0], {}).get(open_set[1]) != "Jam session":
            bad.append("text typed into a set wasn't saved")
        sf_.save_changes(sdir, result.combos, typed={open_set: ""})
        if sf_.load(sdir, result.combos, s).typed:
            bad.append("clearing the text didn't open the set again")
    try:
        sf_.save_changes(sdir, result.combos, sets={(d1, 99): c1})
        bad.append("a set that isn't in the schedule was accepted")
    except sf_.ScheduleFileError:
        pass
    fewer = {cid: c for cid, c in result.combos.items() if cid != c1}
    if not any(c1 in p_ for p_ in sf_.load(sdir, fewer, s).problems):
        bad.append("a combo that isn't accepted any more isn't reported")
    try:
        sf_.load(sdir, result.combos, replace(s, semester_name="Winter 2099"))
        bad.append("another semester's schedule was loaded")
    except sf_.ScheduleFileError:
        pass
    sch = sf_.load(sdir, result.combos, s)
    write_schedule_xlsx(sdir / "Schedule.xlsx", sch, result.combos, s, lambda e: e, {})
    wb = open_wb(sdir / "Schedule.xlsx")
    if "Changes" not in wb.sheetnames or len(list(wb["Changes"].iter_rows())) < 4:
        bad.append("the export has no Changes sheet with the history")
    if wb.sheetnames[:2] != ["By night", "All sets"] or is_old_schedule(sdir / "Schedule.xlsx") \
            or sf_.load(sdir, result.combos, s).sets != sch.sets:
        bad.append(f"the export: sheets {wb.sheetnames}, or it's taken for an old schedule")
    moved = sf_.archive_semester(sdir, s.semester_name)
    if sf_.has_schedule(sdir) or not (moved / "schedule.json").exists() or not (moved / "Schedule backups").exists():
        bad.append("archiving left the schedule behind")
    # an old, hand-editable Schedule.xlsx: converted once (typed text, supervised nights, a typo reported)
    old = Workbook()
    ws = old.active
    ws.title = "Schedule"
    ws.append(["Date", "Day", "Venue", "Set", "Combo", "Supervised"])
    n0 = result.nights[0]
    names = [result.combos[c].name for c in sorted(result.combos)[:2]]
    ws.append([n0.date, "", n0.venue, 1, names[0], "Yes"])
    ws.append([n0.date, "", n0.venue, 2, "Jam session", "Yes"])
    ws.append([n0.date, "", n0.venue, 3, "Combo 999", "Yes"])
    old.properties.subject = s.semester_name
    old.save(sdir / "Schedule.xlsx")
    sch = sf_.load(sdir, result.combos, s)
    if not sch.converted or sch.sets[n0.date].get(1) != sorted(result.combos)[0] \
            or sch.typed.get(n0.date, {}).get(2) != "Jam session" or sch.supervised != {n0.date}:
        bad.append(f"the old Schedule.xlsx wasn't converted right: {sch.sets.get(n0.date)}, {sch.typed}")
    if (sdir / "Schedule.xlsx").exists() or not any(sf_.schedule_backups(sdir).glob("Schedule (old*.xlsx")):
        bad.append("the old Schedule.xlsx wasn't moved into Schedule backups")
    if not any("Combo 999" in t for _, t in sch.report):
        bad.append("the old file's typo isn't in the report")
    if sf_.load(sdir, result.combos, s).converted:
        bad.append("converted twice")
    ok = not bad
    print(f"{'PASS' if ok else 'FAIL'}  schedule file: saved and read back, changes backed up, text in a set, gone "
          "combos reported, export, archive, old Schedule.xlsx converted")
    for b_ in bad[:6]:
        print("      -", b_)
    failures += not ok

    # Combo numbers never change: withdraw one combo and add a new one; the others keep their numbers.
    from openpyxl import load_workbook
    (app_data(tmp) / "scheduler_data.json").unlink(missing_ok=True)
    make(tmp / "a.xlsx", tmp / "f.xlsx", nights, 33, seed=1, sem=base.semester_name)
    before = {c.ref: c.name for c in load_input(tmp, base, "a.xlsx", "f.xlsx").combos}
    wb = load_workbook(tmp / "a.xlsx")
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    first_ref = sorted(before, key=lambda r: before[r])[0]
    for row in ws.iter_rows(min_row=2):
        if str(row[head.index("Response Id")].value) == first_ref:
            row[head.index("Status")].value = "Withdrawn"
    ws.append([999, "x", base.semester_name, "new.person@mail.mcgill.ca", "new.person@mail.mcgill.ca\nother.one@mail.mcgill.ca",
               "prof@mcgill.ca", "No", "Accepted", None, None, None])
    ws.tables["Approvals"].ref = f"A1:K{ws.max_row}"
    wb.save(tmp / "a.xlsx")
    notes_inp = load_input(tmp, base, "a.xlsx", "f.xlsx")
    after = {c.ref: c.name for c in notes_inp.combos}
    moved = [r for r in after if r in before and after[r] != before[r]]
    gone_warned = any("No longer accepted" in t and before[first_ref] in t for _, t in notes_inp.notes)
    ok = (not moved and first_ref not in after and after.get("999") == f"Combo {len(before) + 1:02d}" and gone_warned)
    print(f"{'PASS' if ok else 'FAIL'}  combo numbers stay put: withdrawn {before[first_ref]} leaves a gap, the new "
          f"combo is {after.get('999')}")
    failures += not ok

    # Email fixes: a typo in the approvals hides that member's conflicts; fixing it in the store brings them back.
    from store import Store
    (app_data(tmp) / "scheduler_data.json").unlink(missing_ok=True)
    make(tmp / "a.xlsx", tmp / "f.xlsx", nights, 33, seed=1, sem=base.semester_name)
    clean = load_input(tmp, base, "a.xlsx", "f.xlsx")
    who = next(e for e in sorted(clean.blocked) if any(e in c.members for c in clean.combos) and e.endswith("@mail.mcgill.ca"))
    typo = who.replace("@", "x@")
    wb = load_workbook(tmp / "a.xlsx")
    for row in wb.worksheets[0].iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str) and who in c.value:
                c.value = c.value.replace(who, typo)
    wb.save(tmp / "a.xlsx")
    broken = load_input(tmp, base, "a.xlsx", "f.xlsx")
    st = Store(tmp)
    st.set_email(typo, who)
    st.save()
    fixed = load_input(tmp, base, "a.xlsx", "f.xlsx")
    members = lambda inp: {e for c in inp.combos for e in c.members}
    ok = (typo in members(broken) and who not in members(broken) and who in members(fixed)
          and typo not in members(fixed) and {c.name: c.members for c in fixed.combos} == {c.name: c.members for c in clean.combos})
    print(f"{'PASS' if ok else 'FAIL'}  email fix: a typo'd address in the approvals is corrected everywhere again")
    failures += not ok

    # A cap too small to fit every combo must stop with a clear message, not a vague solver failure.
    tight = replace(base, extra_slot_policy="open", max_supervised_nights=8)       # 8 nights x 4 sets < 33 combos
    try:
        run_schedule(inp, tight)
        print("FAIL  max_supervised_nights too small was accepted")
        failures += 1
    except ScheduleError as e:
        ok = "supervised nights" in str(e) and "Max supervised nights" in str(e)
        print(f"{'PASS' if ok else 'FAIL'}  too-small supervision cap is explained")
        failures += not ok

    # First-year combos' first show on a supervised night; supervision_timing pulls the nights to one end.
    from core.stats import schedule_stats
    runs, bad = {}, []
    for timing in ("early", "late"):
        s = replace(base, extra_slot_policy="open", supervision_timing=timing)
        r = runs[timing] = run_schedule(inp, s)
        fy = [c for c in r.combos if r.combos[c].first_year]
        missed = [c for c in fy if min(d for d, cs in r.lineup.items() if c in cs) not in r.supervised]
        if not fy or missed or r.stats["first_year_unsupervised"]:
            bad.append(f"{timing}: first-year combos without a supervised first show: {missed} (of {len(fy)})")
        sets = {d: {k: c for k, c in enumerate(cs, start=1)} for d, cs in r.lineup.items()}
        sections, problems = schedule_stats(sets, r.nights, r.combos, inp, s, r.supervised)
        lines = dict(sections)["Supervision"]
        if problems or not any(l.startswith(f"First-year combos whose first show is supervised: {len(fy)} of") for l in lines):
            bad.append(f"{timing}: rule check {problems[:2]} / {lines}")
    mean = {t: sum(d.toordinal() for d in r.supervised) / len(r.supervised) for t, r in runs.items()}
    if not mean["early"] < mean["late"]:
        bad.append(f"supervised nights earlier with 'late' than with 'early' ({mean})")
    ok = not bad
    print(f"{'PASS' if ok else 'FAIL'}  first-year first show supervised; supervised nights earlier / later on request")
    for b_ in bad:
        print("      -", b_)
    failures += not ok

    # Supervised nights per combo: 2 each when asked (the cap raised so it fits).
    two = replace(base, extra_slot_policy="auto", min_supervised_per_combo=2, max_supervised_nights=None)
    res2 = run_schedule(inp, two)
    per = {c: sum(1 for d, cs in res2.lineup.items() if c in cs and d in res2.supervised) for c in res2.combos}
    ok = min(per.values()) >= 2 and not check(res2, inp, two)
    print(f"{'PASS' if ok else 'FAIL'}  2 supervised nights per combo: every combo on at least {min(per.values())}")
    failures += not ok

    # Switches: supervision off = no supervised nights and nothing about them in the rule check; first-year off =
    # no combo is first-year, whatever the approvals say.
    from core.stats import schedule_stats
    (app_data(tmp) / "scheduler_data.json").unlink(missing_ok=True)
    make(tmp / "a.xlsx", tmp / "f.xlsx", nights, 33, seed=1, sem=base.semester_name)
    inp = load_input(tmp, base, "a.xlsx", "f.xlsx")
    s = replace(base, extra_slot_policy="open", every_combo_supervised=False)
    result = run_schedule(inp, s)
    sets = {d: {k: c for k, c in enumerate(cs, start=1)} for d, cs in result.lineup.items()}
    _, problems = schedule_stats(sets, result.nights, result.combos, inp, s, None)
    ok = not result.supervised and not check(result, inp, s) and not problems
    print(f"{'PASS' if ok else 'FAIL'}  supervision off: {len(result.supervised)} supervised nights, rule check clean")
    failures += not ok
    on, _ = validate({**DEFAULTS, "first_year_earliest_date": "2027-02-01"})
    off, _ = validate({**DEFAULTS, "first_year_earliest_date": "2027-02-01", "use_first_year": False})
    fy_on = sum(c.first_year for c in load_input(tmp, on, "a.xlsx", "f.xlsx").combos)
    fy_off = sum(c.first_year for c in load_input(tmp, off, "a.xlsx", "f.xlsx").combos)
    result = run_schedule(load_input(tmp, off, "a.xlsx", "f.xlsx"), replace(off, solver_time_limit_sec=10))
    try:                                                # on without a date: refused, with a plain message
        validate({**DEFAULTS, "use_first_year": True, "first_year_earliest_date": None})
        needs_date = False
    except Exception as e:
        needs_date = "turn first-year combos off" in str(e)
    legacy_off, _ = validate({k: v for k, v in DEFAULTS.items() if k != "use_first_year"} | {"first_year_earliest_date": None})
    ok = (fy_on > 0 and fy_off == 0 and off.first_year_earliest_date is None and not result.stats["early_first_year"]
          and needs_date and not legacy_off.use_first_year)
    print(f"{'PASS' if ok else 'FAIL'}  first-year off: {fy_on} first-year combos with it on, {fy_off} with it off")
    failures += not ok

    # Email rules from the settings: domain fixes, the student domain, and old settings keeping McGill's rules.
    from inputs import EmailRules
    from settings_file import parse_domain_fixes
    bad = []
    if parse_domain_fixes("mcgill.ca -> mail.mcgill.ca; @Gmial.com = gmail.com\n") != {
            "mcgill.ca": "mail.mcgill.ca", "gmial.com": "gmail.com"}:
        bad.append("parse_domain_fixes misread a list")
    for wrong in ("mcgill.ca", "mcgill.ca -> ", "a -> b"):
        try:
            parse_domain_fixes(wrong)
            bad.append(f"parse_domain_fixes accepted {wrong!r}")
        except ValueError:
            pass
    rules = EmailRules("mail.mcgill.ca", {"gmial.com": "gmail.com"})
    if rules.norm(" Ana.Ruiz@GMIAL.com ") != "ana.ruiz@gmail.com" or rules.norm("prof@mcgill.ca") != "prof@mcgill.ca":
        bad.append("EmailRules.norm")
    legacy = {k: v for k, v in DEFAULTS.items() if k not in ("student_email_domain", "professor_email_domain")}
    old, _ = validate(legacy)
    if old.student_email_domain != "mail.mcgill.ca" or old.professor_email_domain != "mcgill.ca" \
            or old.email_domain_fixes:
        bad.append("settings saved before the email settings existed lost McGill's domains (or got domain fixes)")
    kept, _ = validate({**DEFAULTS, "email_domain_fixes": "mcgill.ca -> mail.mcgill.ca"})
    if kept.email_domain_fixes != {"mcgill.ca": "mail.mcgill.ca"}:
        bad.append("domain fixes an older settings file has were dropped")
    notes_ = []                                      # a student on the professors' domain: pointed out
    from core.model import Combo as C_
    from inputs import check_addresses as ca_
    ca_([C_("Combo 01", "Combo 01", frozenset({"ana.ruiz@mcgill.ca", "ben.li@mail.mcgill.ca"}))], {},
        EmailRules("mail.mcgill.ca", prof_domain="mcgill.ca"), notes_)
    if not any("ana.ruiz@mcgill.ca" in t for _, t in notes_):
        bad.append("a member on the professors' domain isn't pointed out")
    try:
        validate({**DEFAULTS, "email_domain_fixes": "nonsense"})
        bad.append("validate accepted unreadable email_domain_fixes")
    except Exception:
        pass
    outside = lambda inp: [t for _, t in inp.notes if "with an email outside" in t]
    anywhere, _ = validate({**DEFAULTS, "student_email_domain": ""})
    elsewhere, _ = validate({**DEFAULTS, "student_email_domain": "school.edu"})
    if outside(load_input(tmp, anywhere, "a.xlsx", "f.xlsx")):
        bad.append("a blank student domain still lists outside members")
    if not outside(load_input(tmp, elsewhere, "a.xlsx", "f.xlsx")):
        bad.append("members outside student_email_domain aren't listed")
    ok = not bad
    print(f"{'PASS' if ok else 'FAIL'}  email rules come from the settings (domain fixes, student domain, old settings)")
    for b in bad:
        print("      -", b)
    failures += not ok

    # Input files from anywhere: picked per data folder, remembered, and back to the data folder's file on reset.
    import app_config
    app_config.CONFIG = tmp / "config.json"
    elsewhere_dir = Path(tempfile.mkdtemp())
    (tmp / "a.xlsx").replace(elsewhere_dir / "renamed approvals.xlsx")
    usual = app_config.input_files(tmp)
    app_config.pick_input_file(tmp, "approvals", elsewhere_dir / "renamed approvals.xlsx")
    files = app_config.input_files(tmp)
    picked_ok = (files["approvals"] == elsewhere_dir / "renamed approvals.xlsx" and files["conflicts"] == usual["conflicts"]
                 and app_config.input_files(elsewhere_dir)["approvals"] == elsewhere_dir / "Approvals.xlsx")
    read_ok = len(load_input(tmp, base, files["approvals"], "f.xlsx").combos) == len(inp.combos)
    app_config.pick_input_file(tmp, "approvals", None)
    ok = picked_ok and read_ok and app_config.input_files(tmp) == usual and usual["approvals"] == tmp / "Approvals.xlsx"
    print(f"{'PASS' if ok else 'FAIL'}  input files: a renamed file in another folder is picked, read, and reset")
    failures += not ok

    # Members changed in the app: added, removed, a new liaison; the approvals file itself never changes.
    from core.checks import analyze
    from store import Store
    (app_data(tmp) / "scheduler_data.json").unlink(missing_ok=True)
    make(tmp / "a.xlsx", tmp / "f.xlsx", nights, 33, seed=1, sem=base.semester_name)
    clean = load_input(tmp, base, "a.xlsx", "f.xlsx")
    sem, a, b = base.semester_name, clean.combos[0], clean.combos[1]
    leaving, newbie = a.liaison, "new.player@mail.mcgill.ca"
    stay = sorted(a.members - {leaving})[0]
    joiner = next(e for e in sorted(clean.blocked, key=lambda e: -len(clean.blocked[e])) if e not in b.members)
    st = Store(tmp)
    st.add_member(sem, a.ref, newbie)
    st.remove_member(sem, a.ref, leaving)
    st.set_liaison(sem, a.ref, stay)
    st.add_member(sem, b.ref, joiner)
    st.save()
    approvals_before = (tmp / "a.xlsx").read_bytes()
    edited = load_input(tmp, base, "a.xlsx", "f.xlsx")
    by_ref = {c.ref: c for c in edited.combos}
    allowed_clean, _ = analyze(clean, base, nights)
    allowed_edit, _ = analyze(edited, base, nights)
    bad = []
    if not (newbie in by_ref[a.ref].members and leaving not in by_ref[a.ref].members):
        bad.append("add/remove not applied")
    if by_ref[a.ref].liaison != stay:
        bad.append("chosen liaison not used")
    if joiner not in by_ref[b.ref].members or allowed_edit[b.id] != allowed_clean[b.id] - clean.blocked[joiner]:
        bad.append("an added member's conflicts don't count for their new combo")
    if (tmp / "a.xlsx").read_bytes() != approvals_before:
        bad.append("the approvals file was changed")
    if not any("changed in the app" in t and a.name in t for _, t in edited.notes):
        bad.append("the check doesn't list the changes")
    st = Store(tmp)                                    # undo everything: back to exactly the approvals
    st.remove_member(sem, a.ref, newbie)
    st.add_member(sem, a.ref, leaving)
    st.set_liaison(sem, a.ref, "")
    st.remove_member(sem, b.ref, joiner)
    st.save()
    back = {c.ref: c for c in load_input(tmp, base, "a.xlsx", "f.xlsx").combos}
    if (back[a.ref].members, back[a.ref].liaison, back[b.ref].members) != (a.members, a.liaison, b.members):
        bad.append("undoing the changes doesn't restore the combos")
    if Store(tmp).data["members"].get(sem):
        bad.append("undone changes are still stored")
    st = Store(tmp)                                    # liaison removed without a new one: the next member stands in
    st.remove_member(sem, a.ref, leaving)
    st.save()
    stand_in = {c.ref: c for c in load_input(tmp, base, "a.xlsx", "f.xlsx").combos}[a.ref]
    if not stand_in.liaison or stand_in.liaison not in stand_in.members:
        bad.append("a combo whose liaison was removed has no liaison")
    st = Store(tmp)                                    # everyone removed: the combo stays, with a warning
    for e in a.members:
        st.remove_member(sem, a.ref, e)
    st.save()
    empty = load_input(tmp, base, "a.xlsx", "f.xlsx")
    _, report = analyze(empty, base, nights)
    if a.ref not in {c.ref for c in empty.combos} or not any(f"{a.name} has no members" in t for _, t in report):
        bad.append("a combo with everyone removed disappears or isn't flagged")
    (app_data(tmp) / "scheduler_data.json").unlink()
    _, report4 = analyze(clean, replace(base, min_members_per_combo=4), nights)
    _, report_off = analyze(clean, replace(base, min_members_per_combo=None), nights)
    small = [c for c in clean.combos if len(c.members) < 4]
    if small and not any("fewer than 4" in t for _, t in report4) or any("fewer than" in t for _, t in report_off):
        bad.append("the members-per-combo warning doesn't follow the setting")
    ok = not bad
    (app_data(tmp) / "scheduler_data.json").unlink(missing_ok=True)   # withdrawn in the app: a gap, never reused, put back
    first = {c.ref: c.name for c in load_input(tmp, base, "a.xlsx", "f.xlsx").combos}
    st = Store(tmp)
    st.set_withdrawn(sem, a.ref, True)
    st.save()
    wd = load_input(tmp, base, "a.xlsx", "f.xlsx")
    if a.ref in {c.ref for c in wd.combos} or [c.name for c in wd.withdrawn] != [first[a.ref]]:
        bad.append("a withdrawn combo is still scheduled, or isn't listed as withdrawn")
    if {c.ref: c.name for c in wd.combos} != {r: n for r, n in first.items() if r != a.ref}:
        bad.append("withdrawing a combo renumbered the others")
    if any("No longer accepted" in t_ for _, t_ in wd.notes):
        bad.append("an app-withdrawn combo gets the 'no longer accepted' warning")
    st = Store(tmp)
    st.set_withdrawn(sem, a.ref, False)
    st.save()
    if {c.ref: c.name for c in load_input(tmp, base, "a.xlsx", "f.xlsx").combos} != first:
        bad.append("putting a combo back doesn't give it its old number")
    (app_data(tmp) / "scheduler_data.json").unlink(missing_ok=True)   # first-year tag set in the app wins; off = never
    plain = {c.ref: c for c in load_input(tmp, base, "a.xlsx", "f.xlsx").combos}
    fy, not_fy = next(c for c in plain.values() if c.first_year), next(c for c in plain.values() if not c.first_year)
    st = Store(tmp)
    st.set_first_year(sem, fy.ref, False)
    st.set_first_year(sem, not_fy.ref, True)
    st.save()
    tagged = {c.ref: c for c in load_input(tmp, base, "a.xlsx", "f.xlsx").combos}
    off_fy, _ = validate({**DEFAULTS, "use_first_year": False})
    if tagged[fy.ref].first_year or not tagged[not_fy.ref].first_year:
        bad.append("the first-year tag set in the app doesn't win over the approvals")
    if any(c.first_year for c in load_input(tmp, off_fy, "a.xlsx", "f.xlsx").combos):
        bad.append("a first-year tag set in the app applies while first-year combos are off")
    st = Store(tmp)
    st.set_first_year(sem, fy.ref, None)
    st.set_first_year(sem, not_fy.ref, None)
    st.save()
    cleared = {c.ref: c.first_year for c in load_input(tmp, base, "a.xlsx", "f.xlsx").combos}
    if cleared != {r: c.first_year for r, c in plain.items()}:
        bad.append("clearing the app's first-year tag doesn't go back to the approvals")
    pend = load_input(tmp, base, "a.xlsx", "f.xlsx").pending      # waiting for a decision: listed, not scheduled
    if not pend or any(not c.members for c in pend) or {c.ref for c in pend} & set(plain):
        bad.append("pending combos aren't listed with their members, or are scheduled")
    (tmp / "half.xlsx").write_bytes((tmp / "a.xlsx").read_bytes()[:3000])      # caught mid-sync
    try:
        load_input(tmp, base, "half.xlsx", "f.xlsx")
        bad.append("a half-written approvals file was read")
    except Exception as e:
        if "middle of syncing" not in str(e):
            bad.append(f"a half-written approvals file gives {type(e).__name__}, not a plain message")
    ok = not bad
    print(f"{'PASS' if ok else 'FAIL'}  member changes in the app: add, remove, liaison, undo; approvals untouched; half-synced file explained")
    for b_ in bad:
        print("      -", b_)
    failures += not ok

    # several computers on one data folder: stale saves refused, the lock file, sync apps' conflict copies
    import json
    import time
    import shared_folder as sf
    bad = []
    a, b = Store(tmp), Store(tmp)                          # two computers read scheduler_data.json
    a.set_name("x@mail.mcgill.ca", "From A")
    a.save()
    b.set_name("y@mail.mcgill.ca", "From B")
    try:
        b.save()
        bad.append("a save over a file changed since it was read went through")
    except sf.ChangedOnDisk:
        pass
    if Store(tmp).names.get("x@mail.mcgill.ca") != "From A":
        bad.append("the refused save damaged the other computer's change")
    a.set_name("x@mail.mcgill.ca", "Again A")              # saving twice from one Store is fine
    a.save()
    lock_dir = tmp / "lockdir"
    lock_dir.mkdir()
    if sf.holder(lock_dir):
        bad.append("an empty folder is reported in use")
    sf.claim(lock_dir)
    if sf.holder(lock_dir) or sf.refresh(lock_dir):
        bad.append("our own lock is reported as someone else's")
    other = {"computer": "OFFICE-PC", "user": "ana", "pid": 1, "since": time.time(), "seen": time.time()}
    lock_path(lock_dir).write_text(json.dumps(other))
    if not sf.holder(lock_dir) or "OFFICE-PC (ana)" not in sf.describe(sf.holder(lock_dir)):
        bad.append("another computer's lock isn't reported")
    if not sf.refresh(lock_dir) or sf.read_lock(lock_dir)["computer"] != "OFFICE-PC":
        bad.append("a refresh overwrote the lock of a computer that took the folder over")
    sf.release(lock_dir)
    if not lock_path(lock_dir).exists():
        bad.append("closing removed another computer's lock")
    lock_path(lock_dir).write_text(json.dumps({**other, "seen": time.time() - sf.STALE - 60}))
    if sf.holder(lock_dir):
        bad.append("a left-over lock (not updated for a long time) still counts")
    sf.claim(lock_dir)
    sf.release(lock_dir)
    if lock_path(lock_dir).exists():
        bad.append("closing didn't remove our lock")
    for n in ("Schedule-OFFICE-PC.xlsx", "Conflicts (1).xlsx", "Schedule.xlsx", "Schedule notes.txt",
              "Combos-OFFICE-PC-2.pdf", "Schedule (LAPTOP's conflicted copy 2026-10-06).pdf", "Schedule (draft).pdf"):
        (lock_dir / n).write_text("x")
    (app_data(lock_dir) / "settings-LAPTOP.json").write_text("{}")
    (app_data(lock_dir) / "settings.json.bak").write_text("{}")
    copies = sorted(c.name for c, _ in sf.conflict_copies(lock_dir))
    if copies != ["Combos-OFFICE-PC-2.pdf", "Conflicts (1).xlsx", "Schedule (LAPTOP's conflicted copy 2026-10-06).pdf",
                  "Schedule-OFFICE-PC.xlsx", "settings-LAPTOP.json"]:
        bad.append(f"conflict copies found: {copies}")
    ok = not bad
    print(f"{'PASS' if ok else 'FAIL'}  shared folder: stale saves refused, lock file taken over / left over / released, conflict copies found")
    for b_ in bad:
        print("      -", b_)
    failures += not ok
    # The data folder can be anywhere: usable or not (and why), and a backup restored into a new folder holds
    # everything that matters, with an input file picked from outside the folder under its usual name.
    import os
    import zipfile
    import backup
    import data_folder as dfo
    bad = []
    nowhere, a_file = tmp / "not there", tmp / "a file.txt"
    a_file.write_text("x")
    if not dfo.problem(nowhere) or not dfo.problem(a_file) or dfo.problem(tmp):
        bad.append(f"usable folders: {dfo.problem(nowhere)!r}, {dfo.problem(a_file)!r}, {dfo.problem(tmp)!r}")
    if hasattr(os, "geteuid") and os.geteuid() != 0:      # root can write anywhere
        locked = tmp / "read-only"
        locked.mkdir()
        locked.chmod(0o500)
        if not dfo.problem(locked):
            bad.append("a folder that can't be written to passes as usable")
        locked.chmod(0o700)
    other_dir = tmp / "some other folder"
    other_dir.mkdir()
    (other_dir / "holiday.jpg").write_text("x")
    src = tmp / "to back up"
    src.mkdir()
    if dfo.looks_like_data_folder(other_dir) or not dfo.looks_like_data_folder(tmp / "lockdir") \
            or not dfo.looks_like_data_folder(src):
        bad.append("looks_like_data_folder: only a new (empty) folder or one with the scheduler's files is one")
    (src / dfo.SCHEDULE_XLSX).write_text("schedule")
    (src / dfo.ARCHIVE / "Fall 2026").mkdir(parents=True)
    (src / dfo.ARCHIVE / "Fall 2026" / dfo.SCHEDULE_XLSX).write_text("old schedule")
    dfo.schedule_backups(src).mkdir(parents=True)
    (dfo.schedule_backups(src) / "Schedule 2026-10-01 120000.xlsx").write_text("before a swap")
    dfo.settings_path(src).write_text("{}")
    (dfo.app_data(src) / (dfo.SETTINGS_FILE + ".bak")).write_text("{}")
    dfo.store_path(src).write_text("{}")
    dfo.lock_path(src).write_text("{}")
    (src / ".Schedule.xlsx.123.tmp").write_text("half")
    picked = elsewhere_dir / "Approvals from the flow.xlsx"
    picked.write_text("approvals")
    zip_path = src / backup.backup_name()                  # saved inside the folder: left out of the next one
    n = backup.create_backup(src, zip_path, {"approvals": picked, "conflicts": src / dfo.CONFLICTS_FILE})
    backup.create_backup(src, src / "second.zip")
    with zipfile.ZipFile(src / "second.zip") as zf:
        if any(name.startswith(backup.PREFIX) for name in zf.namelist()):
            bad.append("a backup saved in the data folder went into the next backup")
    info = backup.read_backup(zip_path)
    if n != 7 or info["files"] != 7 or info.get("inputs from", {}).get("approvals") != str(picked):
        bad.append(f"backup: {n} files, info {info}")
    restored = backup.restore_backup(zip_path, backup.new_folder(tmp, "restored"))
    got = sorted(p.relative_to(restored).as_posix() for p in restored.rglob("*") if p.is_file())
    want = sorted(["App data/Schedule backups/Schedule 2026-10-01 120000.xlsx", "App data/scheduler_data.json",
                   "App data/settings.json", "App data/settings.json.bak", "Archive/Fall 2026/Schedule.xlsx",
                   "Approvals.xlsx", "Schedule.xlsx"])
    if got != want or (restored / dfo.APPROVALS_FILE).read_text() != "approvals":
        bad.append(f"restored: {got}")
    try:
        backup.restore_backup(zip_path, restored)
        bad.append("a restore went into a folder that already has things in it")
    except backup.BackupError:
        pass
    # restoring into the data folder in use: what's there is saved first; the forms' spreadsheets, the lock, the logs
    # and saved zips stay; the rest is the backup's; restoring the safety zip undoes it
    live = tmp / "live folder"
    (dfo.app_data(live) / "Logs").mkdir(parents=True)
    (live / dfo.APPROVALS_FILE).write_text("new sign-ups")
    (live / dfo.CONFLICTS_FILE).write_text("new conflicts")
    (live / dfo.COMBOS_PDF).write_text("only here")
    (live / "Combo Scheduler backup old.zip").write_text("a zip saved here")
    dfo.schedule_path(live).write_text("today's schedule")
    dfo.lock_path(live).write_text("our lock")
    (dfo.app_data(live) / "Logs" / "pc.txt").write_text("today's log")
    old = tmp / "old state"
    (old / dfo.ARCHIVE / "Fall 2026").mkdir(parents=True)
    (old / dfo.ARCHIVE / "Fall 2026" / dfo.SCHEDULE_XLSX).write_text("past")
    (old / dfo.APPROVALS_FILE).write_text("old sign-ups")
    dfo.app_data(old).mkdir()
    dfo.schedule_path(old).write_text("yesterday's schedule")
    old_zip = tmp / "old.zip"
    backup.create_backup(old, old_zip)
    safety = backup.restore_in_place(old_zip, live)
    state = {p.relative_to(live).as_posix(): p.read_text() for p in live.rglob("*") if p.is_file()
             and backup.BEFORE_RESTORE not in p.parts}
    want = {"Approvals.xlsx": "new sign-ups", "Conflicts.xlsx": "new conflicts", "Combo Scheduler backup old.zip":
            "a zip saved here", "App data/schedule.json": "yesterday's schedule", "App data/In use.json": "our lock",
            "App data/Logs/pc.txt": "today's log", "Archive/Fall 2026/Schedule.xlsx": "past"}
    if state != want or not safety.exists():
        bad.append(f"restore into the folder in use: {state}")
    backup.restore_in_place(old_zip, live, keep_inputs=False)
    if (live / dfo.APPROVALS_FILE).read_text() != "old sign-ups":
        bad.append("restoring with the backup's spreadsheets kept the current ones")
    backup.restore_in_place(safety, live)
    if dfo.schedule_path(live).read_text() != "today's schedule" or (live / dfo.COMBOS_PDF).read_text() != "only here":
        bad.append("restoring the safety backup didn't undo the restore")
    a_file.write_text("not a zip")
    try:
        backup.read_backup(a_file)
        bad.append("a file that isn't a zip passed as a backup")
    except backup.BackupError:
        pass
    with zipfile.ZipFile(tmp / "evil.zip", "w") as zf:
        zf.writestr(backup.INFO, "{}")
        zf.writestr("../outside.txt", "x")
    try:
        backup.read_backup(tmp / "evil.zip")
        bad.append("a zip with a path outside the folder passed as a backup")
    except backup.BackupError:
        pass
    old = tmp / "old names"                                # a folder from before Approvals.xlsx: still read
    old.mkdir()
    (old / dfo.OLD_APPROVALS_FILE).write_text("x")
    if dfo.usual_inputs(old)["approvals"] != old / dfo.OLD_APPROVALS_FILE:
        bad.append("the old Combo Approvals.xlsx isn't read when it's the only one")
    (old / dfo.APPROVALS_FILE).write_text("x")
    if dfo.usual_inputs(old)["approvals"] != old / dfo.APPROVALS_FILE:
        bad.append("Approvals.xlsx doesn't win over the old name")
    ok = not bad
    print(f"{'PASS' if ok else 'FAIL'}  data folder anywhere: usable or not, backup made, restored into a new folder "
          "and into the folder in use (and undone)")
    for b_ in bad:
        print("      -", b_)
    failures += not ok

    # Combos made in the app and conflicts overruled in the app (Combos tab); the log file.
    import app_log
    bad = []
    (app_data(tmp) / "scheduler_data.json").unlink(missing_ok=True)
    make(tmp / "a.xlsx", tmp / "f.xlsx", nights, 33, seed=1, sem=base.semester_name)
    before = load_input(tmp, base, "a.xlsx", "f.xlsx")
    sem = base.semester_name
    top = max(int(c.name.split()[-1]) for c in before.combos)
    victim = next(e for e in sorted(before.blocked) if before.blocked[e])
    gone_day = sorted(before.blocked[victim])[0]
    st = Store(tmp)
    ref = st.add_combo(sem, ["late.one@mail.mcgill.ca", "late.two@mail.mcgill.ca"], "late.one@mail.mcgill.ca",
                       "prof.late@mcgill.ca", first_year=True)
    st.set_overruled(sem, victim, gone_day)
    st.save()
    after = load_input(tmp, base, "a.xlsx", "f.xlsx")
    new = next((c for c in after.combos if c.ref == ref), None)
    if not new or new.name != f"Combo {top + 1:02d}" or new.liaison != "late.one@mail.mcgill.ca" \
            or not new.first_year or new.professor != "prof.late@mcgill.ca" or len(after.combos) != len(before.combos) + 1:
        bad.append(f"a combo made in the app isn't read right: {new}")
    if {c.name for c in before.combos} - {c.name for c in after.combos}:
        bad.append("a combo made in the app changed the others' numbers")
    if gone_day in after.blocked.get(victim, set()) or not any("overruled in the app" in t for _, t in after.notes):
        bad.append("an overruled conflict still counts, or the check doesn't say so")
    st = Store(tmp)
    st.set_withdrawn(sem, ref, True)
    st.set_overruled(sem, victim, gone_day, False)
    st.save()
    again = load_input(tmp, base, "a.xlsx", "f.xlsx")
    if any(c.ref == ref for c in again.combos) or not any(c.ref == ref for c in again.withdrawn):
        bad.append("withdrawing a combo made in the app didn't work")
    if gone_day not in again.blocked.get(victim, set()):
        bad.append("a conflict counted again isn't back")
    st = Store(tmp)
    st.set_overruled(sem, victim, gone_day)
    st.set_email(victim, "renamed.person@mail.mcgill.ca")
    if st.overruled(sem).get("renamed.person@mail.mcgill.ca") != {gone_day}:
        bad.append("an email fix lost the overruled conflict")
    st = Store(tmp)                                        # a conflict added in the app counts
    free = next(e for c in before.combos for e in sorted(c.members) if not before.blocked.get(e))
    st.set_added_conflict(sem, free, nights[0].date)
    st.save()
    if nights[0].date not in load_input(tmp, base, "a.xlsx", "f.xlsx").blocked.get(free, set()):
        bad.append("a conflict added in the app doesn't count")
    # Check flags addresses that make the scheduler miss someone
    from core.model import Combo
    from inputs import EmailRules, check_addresses, not_emails
    if not_emails("ana.ruiz@mail.mcgill.ca\nTBD\nBen Li <ben.li@mail.mcgill.ca>, Carla Diaz") != ["TBD", "Carla Diaz"]:
        bad.append(f"text that isn't an email: {not_emails('ana.ruiz@mail.mcgill.ca\nTBD, Carla Diaz')}")
    notes_ = []
    check_addresses([Combo("Combo 01", "Combo 01", frozenset({"ana.ruiz@mail.mcgill.ca", "ben.li@mail.mcgil.ca"})),
                     Combo("Combo 02", "Combo 02", frozenset({"ana.ruiz@gmail.com", "dev.patel@mail.mcgill.ca"}))],
                    {"dev.patel@gmail.com": {nights[0].date}, "prof@mcgill.ca": set()},
                    EmailRules("mail.mcgill.ca"), notes_)
    said = " | ".join(t for _, t in notes_)
    if not ("typo of 'mail.mcgill.ca'" in said and "ben.li@mail.mcgil.ca" in said and "two addresses" in said
            and "Conflicts from dev.patel@gmail.com" in said and "prof@mcgill.ca" not in said and len(notes_) == 3):
        bad.append(f"address checks: {said}")
    st = Store(tmp)                                        # an instrument is chosen once, then carried over
    st.set_instrument("Fall 2026", "Combo 03", "kai.drums@mail.mcgill.ca", "Drums")
    st.set_instrument("Fall 2026", "Combo 04", "kai.drums@mail.mcgill.ca", "Vibraphone")   # the latest wins
    later = [Combo("Combo 01", "Combo 01", frozenset({"kai.drums@mail.mcgill.ca", "nobody.yet@mail.mcgill.ca"}))]
    got = st.instruments("Winter 2027", later)
    if got != {("Combo 01", "kai.drums@mail.mcgill.ca"): "Vibraphone"}:
        bad.append(f"instruments not carried over by email: {got}")
    st.set_instrument("Winter 2027", "Combo 01", "kai.drums@mail.mcgill.ca", "")
    if st.instruments("Winter 2027", later):
        bad.append("'No instrument' was filled in with the usual one again")
    st = Store(tmp)                                        # no sheets linked: only the combos made in the app
    st.set_linked("approvals", False)
    st.set_linked("conflicts", False)
    st.save()
    alone = load_input(tmp, base, "a.xlsx", "f.xlsx")
    if alone.combos or [c.ref for c in alone.withdrawn] != [ref] or alone.blocked != {free: {nights[0].date}} \
            or any("Approvals:" in t or "Conflicts row" in t for _, t in alone.notes):
        bad.append(f"unlinked sheets were still read: {[c.ref for c in alone.combos]} active, "
                   f"{[c.ref for c in alone.withdrawn]} withdrawn, {len(alone.blocked)} with conflicts")
    st = Store(tmp)
    st.set_linked("approvals", True)
    st.set_linked("conflicts", True)
    st.save()
    logdir = tmp / "log test"
    logdir.mkdir()
    app_log.set_folder(logdir)
    app_log.LIMIT = 2000
    for i in range(60):
        app_log.write(f"line {i}\nsecond line")
    try:
        raise ValueError("boom")
    except ValueError:
        app_log.error("a test")
    text = app_log.path().read_text()
    if "boom" not in text or "line 0\n" in text or len(text) > 2000 + 400 or not text.startswith("["):
        bad.append("the log isn't written, cut down or doesn't keep errors")
    app_log.set_folder(None)
    app_log.write("nowhere")                                # no folder: nothing, no error
    ok = not bad
    print(f"{'PASS' if ok else 'FAIL'}  combos made in the app (numbered next, edited, withdrawn), conflicts "
          "overruled or added in the app, sheets unlinked, Check flags bad addresses, instruments carried over, "
          "the log")
    for b_ in bad:
        print("      -", b_)
    failures += not ok

    print("\nAll good." if not failures else f"\n{failures} scenario(s) failed.")
    return failures


if __name__ == "__main__":
    raise SystemExit(main())
